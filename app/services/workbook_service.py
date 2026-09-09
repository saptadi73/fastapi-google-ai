"""Template-preserving XLSX round-trip for the executable configuration subset."""

import base64
import copy
import io
import json
import re
from datetime import datetime, timedelta, timezone
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import jwt
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill, Protection
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from app.core.config import ROOT, get_settings
from app.core.exceptions import AppError
from app.schemas.configuration import ConfigurationPatch, ETLConfiguration
from app.services.profiling_service import digest

TEMPLATE = ROOT / "docs/Template_Parameter_Google_Sheet_AI_ETL.xlsx"
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
META = "_ETL_Metadata"
REVIEW = "14 Review"
COLUMN_PARAMETERS = {
    "numeric_precision": 21, "numeric_scale": 22, "varchar_length": 23,
    "date_format": 24, "number_locale": 25, "unit_conversion": 26, "source_timezone": 27, "currency_conversion": 28,
}
# Only these cells are writable. Unsupported parameters and identifiers are signed.
EDITABLE = {
    "02 Struktur Kolom": (5, 104, {4, 10, 11, 12, 13, 17, 20, *COLUMN_PARAMETERS.values()}),
    "03 Aturan Cleansing": (5, 1004, {3, 4, 5, 9}),
    "05 Data Quality": (5, 104, {3, 5, 6, 7, 8, 9, 10, 11, 15, 16}),
    "06 Target Database": (5, 104, {6}),
    "10 Data Product Catalog": (5, 5, {1, 11, 15}),
    "11 Metric Definitions": (5, 204, {1, 2, 6, 7, 16}),
}


def editable(sheet, row, col):
    if sheet == REVIEW:
        return ((2 <= row <= 6 or row == 8) and col == 2) or (10 <= row <= 109 and col in (2, 3))
    start, end, cols = EDITABLE.get(sheet, (0, 0, set()))
    return start <= row <= end and col in cols


def token(claims, audience, minutes):
    return jwt.encode(
        {**claims, "aud": audience, "exp": datetime.now(timezone.utc) + timedelta(minutes=minutes)},
        get_settings().jwt_secret.get_secret_value(),
        algorithm="HS256",
    )


def decode(value, audience):
    try:
        return jwt.decode(
            value, get_settings().jwt_secret.get_secret_value(), algorithms=["HS256"], audience=audience
        )
    except (jwt.PyJWTError, TypeError):
        raise AppError(
            "WORKBOOK_TOKEN_INVALID", "Identitas/preview workbook tidak valid atau kedaluwarsa.", 409
        ) from None


def identities(config, snapshot):
    return {
        "config_id": str(config.id),
        "tenant_id": str(config.tenant_id),
        "revision_no": config.revision_no,
        "config_hash": digest(config.configuration_json),
        "fingerprint": config.based_on_fingerprint,
        "snapshot_hash": snapshot.content_hash,
    }


def readonly_values(book):
    return {
        f"{ws.title}!{cell.coordinate}": [cell.data_type, cell.value]
        for ws in book
        if ws.title != META
        for row in ws
        for cell in row
        if cell.value is not None and not editable(ws.title, cell.row, cell.column)
    }


def put(ws, row, col, value):
    cell = ws.cell(row, col)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Never turn business text into an Excel formula.


def render(config, source, sheet):
    book = load_workbook(TEMPLATE)
    c = ETLConfiguration.model_validate(config.configuration_json).model_dump(mode="json")
    for ws in book:
        if ws.title not in ("00 Petunjuk", "09 Kamus Parameter"):
            for row in ws.iter_rows(min_row=5):
                if row[0].value in ("Prompt Template ID", "Audit Field Minimum"):
                    continue
                for cell in row:
                    if not isinstance(cell, MergedCell):
                        cell.value = None
            ws["A2"] = (
                "Kuning: dapat diedit. Kolom lain belum didukung atau merupakan referensi. Approval hanya melalui aplikasi."
            )
    ws = book["01 Sumber Sheet"]
    for col, val in {
        1: source.source_code,
        2: source.name,
        3: source.spreadsheet_id,
        5: sheet.sheet_name,
        6: sheet.range_a1,
        7: "Ya" if sheet.enabled else "Tidak",
        8: str(source.owner_user_id),
        11: "Viewer",
        12: sheet.header_row,
        13: sheet.data_start_row,
    }.items():
        put(ws, 5, col, val)
    for name, position in COLUMN_PARAMETERS.items():
        put(book["02 Struktur Kolom"], 4, position, name + ("_json" if name in ("unit_conversion", "currency_conversion") else ""))
    rule_row = 5
    for i, column in enumerate(c["columns"], 5):
        ws = book["02 Struktur Kolom"]
        for col, val in {
            1: source.source_code,
            2: column["source_column"],
            3: i - 4,
            4: column["business_name"],
            10: column["target_type"],
            11: "Ya" if column["nullable"] else "Tidak",
            12: "Ya" if column["is_primary_key"] else "Tidak",
            13: "Ya" if column["is_business_key"] else "Tidak",
            17: column["pii_classification"],
            20: column["reason"],
        }.items():
            put(ws, i, col, val)
        for name, position in COLUMN_PARAMETERS.items():
            parameter = column[name]
            put(ws, i, position, json.dumps(parameter, ensure_ascii=False) if name in ("unit_conversion", "currency_conversion") else parameter)
        ws = book["06 Target Database"]
        for col, val in {
            1: f"COL{i - 4}",
            2: source.source_code,
            3: column["source_column"],
            4: "trusted",
            6: column["target_column"],
            16: "Ya",
        }.items():
            put(ws, i, col, val)
        for col, formula in {
            5: f"='{REVIEW}'!B5",
            7: f"='02 Struktur Kolom'!J{i}",
            8: f"='02 Struktur Kolom'!K{i}",
            9: f"='02 Struktur Kolom'!L{i}",
            10: f"='02 Struktur Kolom'!M{i}",
            14: f"='{REVIEW}'!B6",
        }.items():
            ws.cell(i, col, formula)
        for priority, operation in enumerate(column["transformation_codes"], 1):
            for col, val in {3: column["source_column"], 4: priority * 10, 5: operation, 9: "Ya"}.items():
                put(book["03 Aturan Cleansing"], rule_row, col, val)
            rule_row += 1
    put(book["05 Data Quality"], 4, 15, "max_age_days")
    put(book["05 Data Quality"], 4, 16, "default_value_json")
    for i, rule in enumerate(c["data_quality_rules"], 5):
        for col, val in {
            3: rule["column"],
            5: rule["rule"],
            6: json.dumps(rule["value"], ensure_ascii=False),
            7: rule["severity"],
            10: rule["threshold_percent"],
            11: rule["owner"],
            15: rule["max_age_days"],
            16: json.dumps(rule["default_value"], ensure_ascii=False),
            8: rule["action_on_fail"],
            9: "Aktif",
        }.items():
            put(book["05 Data Quality"], i, col, val)
    for col, val in {
        1: c["semantic"]["code"],
        11: json.dumps(c["semantic"]["dimensions"]),
        15: json.dumps(c["semantic"]["allowed_roles"]),
    }.items():
        put(book["10 Data Product Catalog"], 5, col, val)
    for col, row in {2: 2, 4: 3, 8: 4}.items():
        book["10 Data Product Catalog"].cell(5, col, f"='{REVIEW}'!B{row}")
    for i, metric in enumerate(c["semantic"]["metrics"], 5):
        for col, val in {
            1: metric["code"],
            2: metric["label"],
            6: metric["column"],
            7: metric["aggregation"],
            16: "Aktif",
        }.items():
            put(book["11 Metric Definitions"], i, col, val)
    ws = book.create_sheet(REVIEW)
    ws.append(["Parameter", "Nilai", "Status"])
    for row, key in enumerate(
        ("dataset_business_name", "dataset_description", "grain", "target_table", "load_strategy"), 2
    ):
        put(ws, row, 1, key)
        put(ws, row, 2, c[key])
    ws.append(
        [
            "Petunjuk",
            "Jawab pertanyaan, pilih Selesai; upload tetap membutuhkan preview dan approval aplikasi.",
        ]
    )
    put(ws, 8, 1, "append_duplicate_policy")
    put(ws, 8, 2, c.get("append_duplicate_policy"))
    ws.cell(9, 1, "Pertanyaan")
    ws.cell(9, 2, "Jawaban")
    ws.cell(9, 3, "Status")
    for i, question in enumerate(c["unresolved_questions"], 10):
        put(ws, i, 1, question)
        put(ws, i, 3, "Terbuka")
    if len(c["unresolved_questions"]) > 100:
        raise AppError("WORKBOOK_CAPACITY", "Lebih dari 100 pertanyaan; gunakan form sebelum mengekspor.")
    if len(c["semantic"]["metrics"]) > 200:
        raise AppError("WORKBOOK_CAPACITY", "Lebih dari 200 metric; gunakan form.")
    # Remove template validation enums which differ from executable API values.
    for ws in book:
        ws.data_validations.dataValidation = []
        limit = max(ws.max_row, EDITABLE.get(ws.title, (0, 0, set()))[1], 109 if ws.title == REVIEW else 0)
        for row in ws.iter_rows(max_row=limit):
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                allowed = editable(ws.title, cell.row, cell.column)
                cell.protection = Protection(locked=not allowed)
                cell.fill = PatternFill("solid", fgColor="FFF2CC" if allowed else "F1F5F9")
        ws.protection.sheet = True
        ws.freeze_panes = "C5" if ws.title != REVIEW else "B10"
    ws = book[REVIEW]
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 18
    book["00 Petunjuk"]["A2"] = (
        "Draft aplikasi: edit sel kuning. Sheet 04, 07, 08, 12, 13 belum dapat diimport. Status Excel tidak memberi approval."
    )
    book["00 Petunjuk"]["E6"] = "=COUNTIF('03 Aturan Cleansing'!$I$5:$I$1004,\"Ya\")"
    return book


def export_workbook(config, source, sheet, snapshot):
    book = render(config, source, sheet)
    claims = {
        **identities(config, snapshot),
        "readonly_hash": digest(readonly_values(book)),
        "tabs": book.sheetnames,
    }
    ws = book.create_sheet(META)
    ws.append(["format", "etl-review-v1"])
    ws.append(["signature", token(claims, "etl-workbook-export", 60 * 24 * 7)])
    ws.sheet_state = "veryHidden"
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def load_upload(encoded):
    try:
        raw = base64.b64decode(encoded, validate=True)
        if len(raw) > 2_000_000:
            raise ValueError()
        with ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or sum(i.file_size for i in entries) > 30_000_000:
                raise ValueError()
            for entry in entries:
                name = entry.filename.lower()
                if any(x in name for x in ("externallinks", "vbaproject", "embeddings")):
                    raise ValueError()
                if name.endswith(".xml"):
                    content = archive.read(entry)
                    if b"<!DOCTYPE" in content or b"<!ENTITY" in content or b"\x00" in content:
                        raise ValueError()
                    if name.startswith("xl/worksheets/"):
                        for col, row in re.findall(rb'<c\b[^>]*\br="([A-Z]+)([0-9]+)"', content):
                            if len(col) > 2 or (len(col) == 2 and col > b"AF") or int(row) > 2000:
                                raise ValueError()
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        if len(book.sheetnames) > 16:
            raise ValueError()
        # Actual rows are bounded too; worksheet dimension metadata is not trusted.
        for ws in book:
            ws.reset_dimensions()
            count = 0
            for row in ws.iter_rows():
                count += 1
                if count > 2000 or len(row) > 32:
                    raise ValueError()
        book.close()
        return load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    except (ValueError, BadZipFile, KeyError, OSError, TypeError, ParseError, InvalidFileException):
        raise AppError(
            "WORKBOOK_INVALID",
            "Gunakan XLSX hasil export aplikasi, maksimal 2 MB, tanpa macro/link eksternal.",
        ) from None


def parse_workbook(encoded, config, source, sheet, snapshot):
    book = load_upload(encoded)
    if META not in book.sheetnames:
        raise AppError(
            "WORKBOOK_IDENTITY_REQUIRED", "Unduh draft Excel dari aplikasi ini terlebih dahulu.", 409
        )
    claims = decode(book[META]["B2"].value, "etl-workbook-export")
    if any(claims.get(key) != value for key, value in identities(config, snapshot).items()):
        raise AppError(
            "WORKBOOK_STALE",
            "Workbook berasal dari draft, tenant, revision, atau snapshot berbeda; unduh ulang.",
            409,
        )
    if book.sheetnames != [*claims.get("tabs", []), META]:
        raise AppError("WORKBOOK_STRUCTURE", "Jangan menambah, menghapus, atau mengganti nama tab template.")
    locked = readonly_values(book)
    if digest(locked) != claims["readonly_hash"]:
        baseline = readonly_values(render(config, source, sheet))
        changed = [key for key in set(locked) | set(baseline) if locked.get(key) != baseline.get(key)]
        return {
            "configuration": None,
            "question_answers": {},
            "errors": [
                {
                    "location": key,
                    "message": "Parameter belum didukung atau sel referensi tidak boleh diubah.",
                }
                for key in sorted(changed)[:50]
            ]
            or [{"location": "workbook", "message": "Sel referensi telah berubah."}],
        }
    errors = []
    for ws in book:
        for row in ws:
            for cell in row:
                if editable(ws.title, cell.row, cell.column) and cell.data_type in ("f", "e"):
                    errors.append(
                        {
                            "location": f"{ws.title}!{cell.coordinate}",
                            "message": "Isi nilai literal, bukan formula/error Excel.",
                        }
                    )
    if errors:
        return {"configuration": None, "question_answers": {}, "errors": errors[:50]}
    c = copy.deepcopy(config.configuration_json)
    answers = {}

    def value(tab, row, col, default=""):
        v = book[tab].cell(row, col).value
        return default if v is None else v

    def active(tab, row, col, yes, no):
        v = value(tab, row, col)
        if v not in (yes, no):
            raise ValueError(f"{tab}, baris {row}: status harus {yes}/{no}")
        return v == yes

    try:
        for row, key in enumerate(
            ("dataset_business_name", "dataset_description", "grain", "target_table", "load_strategy"), 2
        ):
            c[key] = value(REVIEW, row, 2)
        # Old signed workbooks have no policy row; preserve their configuration.
        if value(REVIEW, 8, 1) == "append_duplicate_policy":
            c["append_duplicate_policy"] = value(REVIEW, 8, 2) or None
        for row, column in enumerate(c["columns"], 5):
            for key, col in {
                "business_name": 4,
                "target_type": 10,
                "pii_classification": 17,
                "reason": 20,
            }.items():
                column[key] = value("02 Struktur Kolom", row, col)
            for key, col in {"nullable": 11, "is_primary_key": 12, "is_business_key": 13}.items():
                column[key] = active("02 Struktur Kolom", row, col, "Ya", "Tidak")
            for name, position in COLUMN_PARAMETERS.items():
                parameter = value("02 Struktur Kolom", row, position)
                column[name] = (json.loads(parameter) if parameter != "" else None) if name in ("unit_conversion", "currency_conversion") else (
                    parameter if parameter != "" else None
                )
            column["target_column"] = value("06 Target Database", row, 6)
            column["transformation_codes"] = []
        # Extra structural rows would otherwise disappear silently.
        for tab, cols in (
            ("02 Struktur Kolom", EDITABLE["02 Struktur Kolom"][2]),
            ("06 Target Database", {6}),
        ):
            for row in range(5 + len(c["columns"]), 105):
                if any(value(tab, row, col) != "" for col in cols):
                    raise ValueError(
                        f"{tab}, baris {row}: penambahan kolom melalui Excel belum didukung; gunakan form"
                    )
        columns = {x["source_column"]: x for x in c["columns"]}
        transforms = {}
        for row in range(5, 1005):
            tab = "03 Aturan Cleansing"
            if not any(value(tab, row, col) != "" for col in (3, 4, 5, 9)):
                continue
            name, priority, operation = value(tab, row, 3), value(tab, row, 4), value(tab, row, 5)
            if (
                name not in columns
                or isinstance(priority, bool)
                or not isinstance(priority, (int, float))
                or priority != int(priority)
            ):
                raise ValueError(f"{tab}, baris {row}: kolom/prioritas tidak valid")
            if active(tab, row, 9, "Ya", "Tidak"):
                if (name, int(priority)) in transforms:
                    raise ValueError(f"{tab}, baris {row}: prioritas kolom duplikat")
                transforms[(name, int(priority))] = operation
        for (name, priority), operation in sorted(transforms.items()):
            columns[name]["transformation_codes"].append(operation)
        c["data_quality_rules"] = []
        for row in range(5, 105):
            tab = "05 Data Quality"
            if not any(value(tab, row, col) != "" for col in (3, 5, 6, 7, 8, 9, 10, 11, 15, 16)):
                continue
            if active(tab, row, 9, "Aktif", "Nonaktif"):
                raw = value(tab, row, 6, "null")
                c["data_quality_rules"].append(
                    {
                        "column": value(tab, row, 3),
                        "rule": value(tab, row, 5),
                        "value": json.loads(raw) if isinstance(raw, str) else raw,
                        "action_on_fail": value(tab, row, 8),
                        "severity": value(tab, row, 7, "ERROR"),
                        "threshold_percent": value(tab, row, 10) if value(tab, row, 10) != "" else None,
                        "owner": value(tab, row, 11) or None,
                        "max_age_days": value(tab, row, 15) if value(tab, row, 15) != "" else None,
                        "default_value": json.loads(str(value(tab, row, 16, "null"))),
                    }
                )
        semantic = c["semantic"]
        semantic["code"] = value("10 Data Product Catalog", 5, 1)
        semantic["dimensions"] = json.loads(value("10 Data Product Catalog", 5, 11))
        semantic["allowed_roles"] = json.loads(value("10 Data Product Catalog", 5, 15))
        semantic["metrics"] = []
        for row in range(5, 205):
            tab = "11 Metric Definitions"
            if not any(value(tab, row, col) != "" for col in (1, 2, 6, 7, 16)):
                continue
            if active(tab, row, 16, "Aktif", "Nonaktif"):
                semantic["metrics"].append(
                    {
                        "code": value(tab, row, 1),
                        "label": value(tab, row, 2),
                        "column": value(tab, row, 6),
                        "aggregation": value(tab, row, 7),
                    }
                )
        remaining = []
        for row, question in enumerate(config.configuration_json.get("unresolved_questions", []), 10):
            if active(REVIEW, row, 3, "Selesai", "Terbuka"):
                answer = value(REVIEW, row, 2)
                if not isinstance(answer, str) or not answer.strip():
                    raise ValueError(f"{REVIEW}, baris {row}: jawaban wajib diisi")
                answers[question] = answer
            else:
                remaining.append(question)
        for row in range(10 + len(config.configuration_json.get("unresolved_questions", [])), 110):
            if value(REVIEW, row, 2) or value(REVIEW, row, 3):
                raise ValueError(f"{REVIEW}, baris {row}: tidak ada pertanyaan untuk jawaban ini")
        c["unresolved_questions"] = remaining
        parsed = ConfigurationPatch(revision_no=config.revision_no, configuration=c, question_answers=answers)
        return {
            "configuration": parsed.configuration.model_dump(mode="json"),
            "question_answers": answers,
            "errors": [],
        }
    except ValidationError as exc:
        errors = [{"location": ".".join(map(str, e["loc"])), "message": e["msg"]} for e in exc.errors()]
    except (ValueError, TypeError) as exc:
        errors = [{"location": "workbook", "message": str(exc)[:500]}]
    return {"configuration": None, "question_answers": {}, "errors": errors}
