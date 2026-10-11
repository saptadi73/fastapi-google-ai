import hashlib
import io
import json
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.schemas.configuration import ETLConfiguration
from app.services.etl_compiler_service import transform_rows
from app.services.help_service import KnowledgeBase
from app.services.profiling_service import profile_values
from scripts.build_analytics_examples import OUTPUT, ROOT, SCENARIOS, build_documents, sheet_collection, value


@pytest.mark.parametrize("scenario", SCENARIOS, ids=lambda scenario: scenario["id"])
def test_excel_cached_values_pass_real_app_profiling_and_typed_validation(scenario):
    workbook = load_workbook(OUTPUT / scenario["filename"], data_only=True)
    sheets = sheet_collection(scenario)
    for spec in sheets:
        ws = workbook[spec.name]
        values = [list(row) for row in ws.iter_rows(values_only=True)]
        profile = profile_values(values, spec.name)
        assert profile["row_count"] == len(spec.rows)
        config = ETLConfiguration.model_validate(spec.config(scenario["id"]))
        good, issues, _ = transform_rows(values, SimpleNamespace(header_row=1, data_start_row=2), config)
        assert not issues and len(good) == len(spec.rows), (spec.name, issues)
        assert not ws.merged_cells.ranges and ws.freeze_panes == "A2"
        assert len({row[spec.key] for row in spec.rows}) == len(spec.rows)
    masters = {spec.key: {row[spec.key] for row in spec.rows} for spec in sheets if spec.role == "MASTER"}
    for spec in sheets:
        for row in spec.rows:
            for key, permitted in masters.items():
                if key in row:
                    assert row[key] in permitted
    source_by_name = {spec.name: spec for spec in sheets}
    for action in source_by_name["Action_Plan"].rows:
        source = source_by_name[action["source_sheet"]]
        referenced = next(row for row in source.rows if row[source.key] == action["source_record_id"])
        assert referenced["period_start"] == action["period_start"]
        assert referenced["period_end"] == action["period_end"]
        assert action["cause_status"] == "UNVERIFIED" and action["followup_value"] is None


def test_arithmetic_and_missing_data_have_correct_business_meaning():
    scenarios = {scenario["id"]: {s.name: s for s in sheet_collection(scenario)} for scenario in SCENARIOS}
    for row in scenarios["kepuasan"]["CSAT_Rekap"].rows:
        assert value(row["received_count"]) == value(row["valid_response_count"]) + value(row["excluded_count"])
        assert value(row["csat_ratio"]) == pytest.approx(value(row["satisfied_count"]) / value(row["valid_response_count"]), abs=1e-6)
        assert value(row["response_rate"]) != value(row["csat_ratio"])
    for row in scenarios["marketing"]["Marketing_Data"].rows:
        if row["marketing_type"] == "ATL":
            assert row["leads_count"] is None and value(row["roas_ratio"]) is None
        else:
            assert row["orders_count"] <= row["qualified_leads_count"] <= row["leads_count"]
            assert value(row["cpl_idr"]) == pytest.approx(row["spend_idr"] / row["leads_count"], abs=1e-6)
    prices = scenarios["competitiveness"]["Harga_Observasi"].rows
    own = next(r for r in prices if r["competitor_code"] == "OWN")
    rival = next(r for r in prices if r["competitor_code"] == "COMP-A")
    assert rival["package_price_idr"] < own["package_price_idr"]
    assert value(rival["normalized_price_idr_per_l"]) > value(own["normalized_price_idr_per_l"])
    for row in scenarios["competitiveness"]["Competitive_Survey"].rows:
        assert row["weight_product"] + row["weight_service"] + row["weight_value"] == pytest.approx(1)
    for row in scenarios["kpi"]["KPI_Result"].rows:
        target, actual = value(row["target_value"]), value(row["actual_value"])
        higher = row["direction"] == "HIGHER_IS_BETTER"
        assert value(row["unfavourable_gap"]) == pytest.approx(max(0, target - actual if higher else actual - target), abs=1e-6)
        assert value(row["achievement_ratio"]) == pytest.approx(actual / target if higher else target / actual, abs=1e-6)


def test_package_reproducible_formula_cache_and_download_integrity():
    expected = build_documents()
    for path, content in expected.items():
        assert path.read_bytes() == content
    catalog = json.loads((ROOT / "docs/analytics-examples.json").read_text(encoding="utf-8"))
    assert len(catalog["examples"]) == 5
    with ZipFile(OUTPUT / "Paket_5_Contoh_Analitik.zip") as package:
        for entry in catalog["examples"]:
            content = package.read(entry["filename"])
            assert hashlib.sha256(content).hexdigest() == entry["sha256"]
            formulas = load_workbook(io.BytesIO(content), data_only=False)
            cached = load_workbook(io.BytesIO(content), data_only=True)
            formula_count = 0
            for sheet in formulas:
                for row in sheet:
                    for cell in row:
                        if cell.data_type == "f":
                            formula_count += 1
                            item = cached[sheet.title][cell.coordinate]
                            assert item.data_type != "e", (sheet.title, cell.coordinate)
                            assert item.value is not None or "IF(" in cell.value
            assert formula_count > 0


@pytest.mark.parametrize("scenario", SCENARIOS, ids=lambda scenario: scenario["id"])
def test_non_technical_helper_finds_example_file_and_limitations(scenario):
    articles = KnowledgeBase().search("contoh excel " + scenario["title"], "TECHNICAL_APPROVER", "/guide")
    article = next(a for a in articles if a.article_id == "contoh-excel-" + scenario["id"])
    assert scenario["filename"] in article.body
    assert scenario["caution"] in article.body
    assert len(article.body) <= 6000
