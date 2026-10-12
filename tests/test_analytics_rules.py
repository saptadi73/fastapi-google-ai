import json
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.exceptions import AppError
from app.schemas.analytics_rules import AnalyticsRuleDefinition, AnalyticsRulesDecision
from app.services.analytics_expression_service import calculate
from app.services.analytics_rules_service import (
    FORMULA_HEADERS,
    RULE_HEADERS,
    AnalyticsRulesService,
    definition_from_tables,
    evaluate_rows,
    semantic_hash,
)
from app.services.query_execution_service import build_query
from scripts.build_analytics_examples import OUTPUT, SCENARIOS, analytics_definition, sheet_collection


def definition():
    return AnalyticsRuleDefinition(
        name="Revenue review",
        product_code="SALES",
        query={
            "metrics": ["actual", "target"],
            "dimensions": ["unit"],
            "filters": [{"field": "period", "operator": "between", "value": ["2026-01-01", "2026-01-31"]}],
        },
        formulas=[{"formula_code": "achievement", "expression": "actual / target", "output_unit": "RATIO"}],
        rules=[
            {
                "rule_code": "pipeline",
                "formula_code": "achievement",
                "operator": "gt",
                "threshold": "0.50",
                "action_description": "Review pipeline",
                "owner_team_code": "TEAM-SALES",
                "due_days": 7,
            }
        ],
    )


@pytest.mark.parametrize(
    "expression",
    [
        "__import__('os')",
        "actual.__class__",
        "actual[0]",
        "2 ** 10000",
        "sum(actual)",
        "actual // target",
        "True",
        "1e100",
        "unknown + 1",
        "0.8",
    ],
)
def test_untrusted_formula_is_rejected(expression):
    payload = definition().model_dump(mode="json")
    payload["formulas"][0]["expression"] = expression
    with pytest.raises(ValidationError):
        AnalyticsRuleDefinition.model_validate(payload)


def test_precise_decimal_null_zero_and_dependency():
    assert calculate("0.100000000000000001 + actual", {"actual": "0.2"}) == (
        Decimal("0.300000000000000001"),
        None,
    )
    assert calculate("actual / target", {"actual": "20", "target": 0}) == (None, "ZERO_DENOMINATOR")
    assert calculate("actual / target", {"actual": None, "target": 10}) == (None, "MISSING_VALUE")
    assert calculate("actual", {"actual": True}) == (None, "NON_NUMERIC_VALUE")
    assert calculate("actual * target", {"actual": "1e25", "target": "1e20"}) == (
        None,
        "NUMERIC_OUT_OF_RANGE",
    )
    payload = definition().model_dump(mode="json")
    payload["formulas"].append(
        {"formula_code": "percentage", "expression": "achievement * 100", "output_unit": "PERCENT"}
    )
    data = AnalyticsRuleDefinition.model_validate(payload)
    rows = evaluate_rows(data, [{"unit": "U1", "actual": "60", "target": "100"}], today=date(2026, 1, 15))
    assert rows[0]["formulas"][1]["value"] == "60.0"
    assert rows[0]["action_plans"][0]["due_date"] == "2026-01-22"
    payload["formulas"].reverse()
    with pytest.raises(ValidationError):
        AnalyticsRuleDefinition.model_validate(payload)


def test_threshold_boundaries_no_missing_actions_and_evidence():
    data = definition()
    result = evaluate_rows(
        data,
        [
            {"unit": "U1", "actual": "50", "target": "100"},
            {"unit": "U2", "actual": "51", "target": "100"},
            {"unit": "U3", "actual": "10", "target": "0"},
        ],
    )
    assert not result[0]["action_plans"]
    assert result[1]["evidence_metrics"] == {"actual": "51", "target": "100"}
    assert result[1]["action_plans"][0]["status"] == "PROPOSED"
    assert result[2]["formulas"][0]["blocker"] == "ZERO_DENOMINATOR"
    assert not result[2]["action_plans"]
    assert evaluate_rows(data, []) == []


def test_helper_explains_rules_workflow_and_api_limitations():
    from app.services.help_service import KnowledgeBase

    articles = KnowledgeBase().search("Formula_Analitik Aturan_Action_Plan approval evaluate", "VIEWER", "/guide")
    article = next(a for a in articles if a.article_id == "aturan-analitik-action-plan")
    assert len(article.body) <= 6000
    assert "Dikecualikan dari ETL" in article.body and "wizard ETL/chat AI belum" in article.body
    assert "ZERO_DENOMINATOR" in article.body and "REVOKE" in article.body


@pytest.mark.parametrize("scenario", SCENARIOS, ids=lambda item: item["id"])
def test_templates_control_tabs_round_trip_and_approved_semantic_query(scenario):
    sheets = sheet_collection(scenario)
    payload = analytics_definition(scenario, sheets)
    data = AnalyticsRuleDefinition.model_validate(payload)
    workbook = load_workbook(OUTPUT / scenario["filename"], data_only=True)
    tables = [
        [list(row) for row in workbook[name].iter_rows(values_only=True)]
        for name in ("Formula_Analitik", "Aturan_Action_Plan")
    ]
    assert tables[0][0] == FORMULA_HEADERS and tables[1][0] == RULE_HEADERS
    restored = definition_from_tables(data, tables)
    assert restored == data
    spec = next(s for s in sheets if s.config(scenario["id"])["semantic"]["code"] == data.product_code)
    config = spec.config(scenario["id"])
    product = SimpleNamespace(**config["semantic"], columns=config["columns"], view_name="v_example")
    user = SimpleNamespace(tenant_id=str(uuid4()), row_scope={})
    build_query(product, user, data.query)
    assert json.loads((OUTPUT / "rules" / f"{scenario['id']}.json").read_text(encoding="utf-8")) == payload


def service_fixture(monkeypatch, *, creator="owner", status="DRAFT"):
    data = definition()
    product = SimpleNamespace(
        id="product",
        version=1,
        source_sheet_id="sheet",
        columns=[
            {"target_column": "period", "target_type": "date"},
            {"target_column": "unit", "target_type": "text"},
            {"target_column": "amount", "target_type": "numeric"},
        ],
        metrics=[{"code": code, "column": "amount", "aggregation": "sum"} for code in ("actual", "target")],
        dimensions=["period", "unit"],
    )
    bundle = SimpleNamespace(
        id="bundle",
        product_id=product.id,
        source_id=None,
        created_by=creator,
        status=status,
        definition=data.model_dump(mode="json"),
        review_hash="a" * 64,
        semantic_hash=semantic_hash(product),
        sheet_hash=None,
    )
    actor = SimpleNamespace(id="reviewer", tenant_id="tenant", role="TECHNICAL_APPROVER")
    service = AnalyticsRulesService(Mock(), actor)
    service.session.refresh = AsyncMock()
    service.repo.get = AsyncMock(
        side_effect=lambda cls, key, **kw: bundle if key == "bundle" else SimpleNamespace(source_id="source")
    )
    service.queries.repo.product = AsyncMock(return_value=product)
    service.queries.compile = AsyncMock(return_value=(product, [], [], None, None))
    service.queries.execute = AsyncMock(
        return_value={"rows": [{"unit": "U1", "actual": "60", "target": "100"}], "meta": {}}
    )
    monkeypatch.setattr("app.services.analytics_rules_service.require_source_approver", AsyncMock())
    monkeypatch.setattr("app.services.analytics_rules_service.record", lambda obj: vars(obj).copy())
    return service, bundle, product


@pytest.mark.asyncio
async def test_review_and_evaluate_require_approval_and_preserve_audit(monkeypatch):
    service, bundle, _ = service_fixture(monkeypatch)
    with pytest.raises(AppError, match="belum disetujui"):
        await service.evaluate("bundle")
    preview = await service.evaluate("bundle", preview=True)
    assert preview["mode"] == "PREVIEW"
    await service.decide(
        "bundle", AnalyticsRulesDecision(expected_hash="a" * 64, decision="APPROVE", reason="Reviewed")
    )
    result = await service.evaluate("bundle")
    assert result["rows"][0]["action_plans"][0]["observed_value"] == "0.6"
    assert bundle.approved_by == "reviewer" and bundle.status == "APPROVED"
    assert service.session.add.call_count >= 3
    service.queries.execute.assert_awaited_with(
        "SALES", definition().query, query_source="ANALYTICS_RULES", preserve_numeric=True
    )
    await service.decide(
        "bundle", AnalyticsRulesDecision(expected_hash="a" * 64, decision="REVOKE", reason="Target changed")
    )
    assert bundle.status == "REVOKED"
    with pytest.raises(AppError, match="belum disetujui"):
        await service.evaluate("bundle")


@pytest.mark.asyncio
@pytest.mark.parametrize("failure", ["self", "hash", "semantic", "sheet", "permission", "truncated"])
async def test_stale_self_denied_and_truncated_fail_closed(monkeypatch, failure):
    service, bundle, product = service_fixture(monkeypatch)
    if failure == "self":
        bundle.created_by = service.user.id
    elif failure == "hash":
        bundle.review_hash = "b" * 64
    elif failure == "semantic":
        product.version = 2
    elif failure == "sheet":
        bundle.source_id = "source"
        bundle.sheet_hash = "a" * 64
        service.read_tables = AsyncMock(return_value=[[FORMULA_HEADERS], [RULE_HEADERS]])
    elif failure == "permission":
        service.queries.repo.product.side_effect = AppError("FORBIDDEN", "No access", 403)
    else:
        bundle.status = "APPROVED"
        service.queries.execute.return_value["rows"] *= definition().query.limit
    with pytest.raises(AppError):
        if failure == "truncated":
            await service.evaluate("bundle")
        else:
            await service.decide(
                "bundle",
                AnalyticsRulesDecision(expected_hash="a" * 64, decision="APPROVE", reason="Reviewed"),
            )
    if failure != "truncated":
        assert bundle.status == "DRAFT"


@pytest.mark.integration
@pytest.mark.asyncio
async def test_new_table_enforces_tenant_links_and_stores_immutable_definition():
    from sqlalchemy.exc import IntegrityError

    from app.core.database import SessionFactory
    from app.models import AnalyticsRuleBundle, DataProduct, DataSource, SourceSheet, Tenant, User

    async with SessionFactory() as session:
        tenant, other = Tenant(code="rules_" + uuid4().hex), Tenant(code="rules_" + uuid4().hex)
        session.add_all([tenant, other])
        await session.flush()
        user = User(tenant_id=tenant.id, username="owner", password_hash="unused")
        session.add(user)
        await session.flush()
        source = DataSource(
            tenant_id=tenant.id, source_code="TEST", name="TEST", spreadsheet_id="TEST", owner_user_id=user.id
        )
        session.add(source)
        await session.flush()
        sheet = SourceSheet(tenant_id=tenant.id, source_id=source.id, sheet_id=1, sheet_name="Data")
        session.add(sheet)
        await session.flush()
        product = DataProduct(
            tenant_id=tenant.id,
            source_sheet_id=sheet.id,
            code="SALES",
            name="Sales",
            view_name="v_test",
            columns=[],
            metrics=[],
            dimensions=[],
            allowed_roles=[],
        )
        session.add(product)
        await session.flush()
        good = AnalyticsRuleBundle(
            tenant_id=tenant.id,
            product_id=product.id,
            created_by=user.id,
            definition=definition().model_dump(mode="json"),
            review_hash="a" * 64,
            semantic_hash="b" * 64,
        )
        session.add(good)
        await session.flush()
        assert good.status == "DRAFT"
        with pytest.raises(IntegrityError):
            async with session.begin_nested():
                session.add(
                    AnalyticsRuleBundle(
                        tenant_id=other.id,
                        product_id=product.id,
                        created_by=user.id,
                        definition={},
                        review_hash="a" * 64,
                        semantic_hash="b" * 64,
                    )
                )
                await session.flush()
        await session.rollback()
