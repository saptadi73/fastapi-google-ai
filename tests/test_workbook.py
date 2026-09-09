import base64
import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.core.exceptions import AppError
from app.schemas.configuration import ETLConfiguration
from app.services.workbook_service import REVIEW, export_workbook, parse_workbook


@pytest.fixture
def workbook_context(config_data):
    config = SimpleNamespace(
        id="config",
        tenant_id="tenant",
        revision_no=1,
        based_on_fingerprint="abc",
        configuration_json=ETLConfiguration(**config_data).model_dump(mode="json"),
    )
    source = SimpleNamespace(source_code="SALES", name="Sales", spreadsheet_id="sheet", owner_user_id="owner")
    sheet = SimpleNamespace(sheet_name="Sales", range_a1="A:CV", enabled=True, header_row=1, data_start_row=2)
    snapshot = SimpleNamespace(content_hash="a" * 64)
    return config, source, sheet, snapshot


def encode(raw):
    return base64.b64encode(raw).decode()


def test_append_policy_roundtrip_and_edit(workbook_context):
    config = workbook_context[0].configuration_json
    config.update(load_strategy="APPEND", append_duplicate_policy="SKIP_IDENTICAL")
    for column in config["columns"]:
        column["is_business_key"] = column["is_primary_key"] = False
    result = parse_workbook(encode(export_workbook(*workbook_context)), *workbook_context)
    assert not result["errors"] and result["configuration"]["append_duplicate_policy"] == "SKIP_IDENTICAL"
    result = parse_workbook(edited(workbook_context, REVIEW, "B8", "REJECT_IDENTICAL"), *workbook_context)
    assert not result["errors"] and result["configuration"]["append_duplicate_policy"] == "REJECT_IDENTICAL"


def edited(ctx, tab, cell, value):
    book = load_workbook(io.BytesIO(export_workbook(*ctx)))
    book[tab][cell] = value
    stream = io.BytesIO()
    book.save(stream)
    return encode(stream.getvalue())


def test_roundtrip_preserves_configuration_and_all_tabs(workbook_context):
    raw = export_workbook(*workbook_context)
    assert len(load_workbook(io.BytesIO(raw)).sheetnames) == 16
    result = parse_workbook(encode(raw), *workbook_context)
    assert result["errors"] == []
    assert result["configuration"] == workbook_context[0].configuration_json


def test_edit_identity_and_reject_unsupported_or_formula(workbook_context):
    result = parse_workbook(edited(workbook_context, REVIEW, "B2", "Penjualan baru"), *workbook_context)
    assert result["errors"] == []
    assert result["configuration"]["dataset_business_name"] == "Penjualan baru"
    for tab, cell, value in [
        ("04 Taxonomy Mapping", "C5", "NEW"),
        (REVIEW, "B2", "=1+1"),
        ("02 Struktur Kolom", "D50", "extra column"),
    ]:
        assert parse_workbook(edited(workbook_context, tab, cell, value), *workbook_context)["errors"]


def test_stale_and_invalid_upload(workbook_context):
    raw = encode(export_workbook(*workbook_context))
    workbook_context[0].revision_no += 1
    with pytest.raises(AppError, match="revision"):
        parse_workbook(raw, *workbook_context)
    with pytest.raises(AppError):
        parse_workbook("not base64", *workbook_context)


def test_question_cannot_be_closed_without_answer(workbook_context):
    workbook_context[0].configuration_json["unresolved_questions"] = ["Apakah ID unik?"]
    assert parse_workbook(edited(workbook_context, REVIEW, "C10", "Selesai"), *workbook_context)["errors"]


def test_unit_and_column_parameters_roundtrip_and_edit(workbook_context):
    column = workbook_context[0].configuration_json["columns"][3]
    column.update(numeric_precision=12, numeric_scale=2, number_locale="ID",
                  transformation_codes=["parse_decimal_id"],
                  unit_conversion={"from_unit": "KG", "to_unit": "G", "factor": "1000",
                                   "output_scale": 2, "rounding": "HALF_UP", "on_error": "REJECT_ROW"})
    workbook_context[0].configuration_json["columns"][2]["varchar_length"] = 30
    workbook_context[0].configuration_json["columns"][1]["date_format"] = "%d/%m/%Y"
    result = parse_workbook(encode(export_workbook(*workbook_context)), *workbook_context)
    assert result["errors"] == []
    assert result["configuration"] == workbook_context[0].configuration_json
    result = parse_workbook(edited(workbook_context, "02 Struktur Kolom", "U8", 14), *workbook_context)
    assert result["errors"] == []
    assert result["configuration"]["columns"][3]["numeric_precision"] == 14
    result = parse_workbook(edited(workbook_context, "02 Struktur Kolom", "Z8", '{"from_unit":"USD"}'), *workbook_context)
    assert result["errors"]


def test_timezone_workbook_roundtrip_and_edit(workbook_context):
    column = workbook_context[0].configuration_json["columns"][1]
    column.update(target_type="timestamptz", source_timezone="Asia/Jakarta", transformation_codes=[])
    result = parse_workbook(encode(export_workbook(*workbook_context)), *workbook_context)
    assert result["errors"] == []
    assert result["configuration"] == workbook_context[0].configuration_json
    result = parse_workbook(edited(workbook_context, "02 Struktur Kolom", "AA6", "America/New_York"), *workbook_context)
    assert result["errors"] == []
    assert result["configuration"]["columns"][1]["source_timezone"] == "America/New_York"
    result = parse_workbook(edited(workbook_context, "02 Struktur Kolom", "AA6", "Not/AZone"), *workbook_context)
    assert result["errors"]


def test_currency_workbook_roundtrip_and_edit(workbook_context):
    import json

    conversion = {"from_currency": "USD", "to_currency": "IDR", "rate": "12345.5",
                  "rate_date": "2026-09-09", "rate_reference": "synthetic-demo-rate",
                  "output_scale": 2, "rounding": "HALF_UP", "on_error": "REJECT_ROW"}
    workbook_context[0].configuration_json["columns"][3]["currency_conversion"] = conversion
    result = parse_workbook(encode(export_workbook(*workbook_context)), *workbook_context)
    assert result["errors"] == []
    assert result["configuration"] == workbook_context[0].configuration_json
    updated = {**conversion, "rate": "12346"}
    result = parse_workbook(edited(workbook_context, "02 Struktur Kolom", "AB8", json.dumps(updated)), *workbook_context)
    assert result["errors"] == []
    assert result["configuration"]["columns"][3]["currency_conversion"] == updated
    updated["rate_reference"] = " "
    result = parse_workbook(edited(workbook_context, "02 Struktur Kolom", "AB8", json.dumps(updated)), *workbook_context)
    assert result["errors"]
