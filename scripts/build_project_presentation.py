"""Build the editable project overview deck from reviewed repository documentation.

Requires python-pptx in the authoring environment only; it is not a runtime dependency.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "PRESENTASI_GOOGLE_SHEET_AI.pptx"

NAVY = RGBColor(12, 28, 41)
DEEP = RGBColor(18, 43, 55)
GREEN = RGBColor(9, 103, 82)
TEAL = RGBColor(17, 135, 122)
MINT = RGBColor(204, 238, 226)
CREAM = RGBColor(247, 249, 245)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(26, 47, 55)
MUTED = RGBColor(88, 108, 113)
LINE = RGBColor(219, 230, 226)
AMBER = RGBColor(207, 139, 58)
PALE_AMBER = RGBColor(252, 239, 215)
PALE_BLUE = RGBColor(225, 238, 242)
RED = RGBColor(165, 68, 70)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def rect(slide, x, y, w, h, fill, radius=False, line=None, lw=0.8):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(lw)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.18
    return shape


def oval(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    return shape


def txt(slide, text, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.MIDDLE, font="Aptos", margin=0.03, spacing=1.08):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for index, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(0)
        p.line_spacing = spacing
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return shape


def rule(slide, x1, y1, x2, y2, color=LINE, width=1.3, arrow=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        from pptx.oxml.xmlchemy import OxmlElement

        ln = shape.line._get_or_add_ln()
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        ln.append(tail)
    return shape


def arrow(slide, x, y, color=TEAL, size=0.24):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(size), Inches(size)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def pill(slide, label, x, y, w, fill=MINT, color=GREEN, size=11):
    rect(slide, x, y, w, 0.36, fill, radius=True)
    txt(slide, label, x + 0.08, y, w - 0.16, 0.36, size, color, True)


def card(slide, x, y, w, h, title, body="", accent=GREEN, fill=WHITE,
         title_size=17, body_size=12):
    rect(slide, x, y, w, h, fill, radius=True, line=LINE)
    rect(slide, x, y + 0.22, 0.06, h - 0.44, accent, radius=True)
    txt(slide, title, x + 0.24, y + 0.2, w - 0.46, 0.52, title_size, INK, True)
    if body:
        txt(slide, body, x + 0.24, y + 0.78, w - 0.48, h - 0.98, body_size, MUTED,
            valign=MSO_ANCHOR.TOP, spacing=1.2)


def page(title, eyebrow, no, dark=False, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, NAVY if dark else CREAM)
    if dark:
        oval(slide, 10.1, -1.2, 4.4, 4.4, DEEP)
        oval(slide, -1.2, 5.8, 3.2, 3.2, DEEP)
    label_color = MINT if dark else GREEN
    title_color = WHITE if dark else INK
    txt(slide, eyebrow.upper(), 0.72, 0.37, 10.5, 0.28, 10.5, label_color, True)
    txt(slide, title, 0.72, 0.78, 11.85, 0.68, 28, title_color, True)
    if subtitle:
        txt(slide, subtitle, 0.75, 1.49, 11.75, 0.48, 13, MINT if dark else MUTED)
    rule(slide, 0.74, 7.1, 12.55, 7.1, RGBColor(77, 104, 104) if dark else LINE)
    txt(slide, "GOOGLE SHEET AI  ·  TINJAUAN APLIKASI", 0.74, 7.15, 7.8, 0.18, 8.5,
        MINT if dark else MUTED)
    txt(slide, f"{no:02d}", 12.0, 7.13, 0.55, 0.21, 9.5, MINT if dark else MUTED,
        align=PP_ALIGN.RIGHT)
    return slide


def notes(slide, text):
    try:
        tf = slide.notes_slide.notes_text_frame
        if tf:
            tf.text = text
    except AttributeError:
        pass


# 01 — Cover
s = page("Google Sheet AI", "Presentasi aplikasi", 1, dark=True)
txt(s, "Dari spreadsheet menjadi data tepercaya,\nanalitik, dan keputusan yang lebih cepat.",
    0.8, 2.0, 8.4, 1.65, 28, WHITE, True, valign=MSO_ANCHOR.TOP)
pill(s, "ETL + GOVERNANCE", 0.83, 4.25, 2.2, TEAL, WHITE)
pill(s, "MASTER + TAXONOMY", 3.16, 4.25, 2.35, TEAL, WHITE)
pill(s, "NL2SQL + CHART", 5.64, 4.25, 2.05, TEAL, WHITE)
txt(s, "Oktober 2026  |  Gambaran bisnis dan proses penggunaan", 0.84, 5.4, 9.5, 0.45,
    15, MINT)
notes(s, "Pembuka: aplikasi mengubah Google Sheet dari sumber kerja harian menjadi data yang divalidasi, diberi konteks bisnis, dan bisa ditanya dengan bahasa alami. Jelaskan bahwa ini gambaran proses, bukan demo teknis API.")

# 02 — Background
s = page("Mengapa aplikasi ini dibutuhkan?", "Latar belakang", 2,
         subtitle="Spreadsheet mudah dipakai, tetapi semakin sulit dipercaya ketika jumlah sumber dan pengguna bertambah.")
for x, title, body in [
    (0.72, "Sumber tersebar", "Banyak file, tab, versi, dan pemilik yang berbeda."),
    (3.88, "Arti data berbeda", "Nama kolom mirip dapat berarti hal yang berlainan."),
    (7.04, "Kontrol manual", "Validasi, izin, dan perubahan sering bergantung pada orang."),
    (10.2, "Insight lambat", "Laporan dan chart harus disusun berulang."),
]:
    card(s, x, 2.3, 2.42, 2.25, title, body, accent=AMBER, title_size=16, body_size=12)
rect(s, 1.1, 5.2, 11.1, 0.83, NAVY, radius=True)
txt(s, "Kebutuhan utama: satu jalur yang menambah konteks, kualitas, persetujuan, dan akses sebelum data dipakai.",
    1.43, 5.32, 10.5, 0.6, 17, WHITE, True, align=PP_ALIGN.CENTER)
notes(s, "Mulai dari masalah yang dialami pengguna: bukan kekurangan spreadsheet, melainkan kurangnya jejak perubahan, makna kolom, dan akses yang konsisten. Hindari klaim bahwa semua file otomatis menjadi tabel baru.")

# 03 — Goals
s = page("Empat hasil yang ingin dicapai", "Tujuan", 3)
goals = [
    (0.8, 2.05, "01", "Data tepercaya", "Profiling, klasifikasi, validasi, preview, dan persetujuan."),
    (6.83, 2.05, "02", "Makna konsisten", "Master, taxonomy, semantic product, metrik, dan dimensi."),
    (0.8, 4.18, "03", "Akses sesuai kewenangan", "Role, unit, domain, yurisdiksi, dan policy yang berlaku."),
    (6.83, 4.18, "04", "Jawaban lebih cepat", "Pertanyaan bahasa alami, tabel, chart, dan bantuan AI."),
]
for x, y, number, title, body in goals:
    rect(s, x, y, 5.7, 1.72, WHITE, radius=True, line=LINE)
    oval(s, x + 0.22, y + 0.24, 0.67, 0.67, MINT)
    txt(s, number, x + 0.28, y + 0.33, 0.56, 0.38, 14, GREEN, True, align=PP_ALIGN.CENTER)
    txt(s, title, x + 1.08, y + 0.22, 4.3, 0.52, 18, INK, True)
    txt(s, body, x + 1.08, y + 0.83, 4.28, 0.55, 13, MUTED)
notes(s, "Tekankan empat hasil yang dirasakan pengguna. Data dianggap tepercaya setelah melewati proses, bukan hanya setelah spreadsheet berhasil dibaca.")

# 04 — prerequisites
s = page("Apa yang perlu disiapkan?", "Syarat awal", 4,
         subtitle="Persiapan dibagi antara pemilik data, pengelola aplikasi, dan pengguna analitik.")
prereqs = [
    (0.8, "01", "Sumber Google Sheet", "URL/ID spreadsheet; akses Viewer untuk service account; tab dan header yang jelas."),
    (4.97, "02", "Konteks bisnis", "Pemilik, steward, unit, domain, yurisdiksi, purpose, dan sensitivitas."),
    (9.14, "03", "Akun & layanan", "Akun ber-role; PostgreSQL dan layanan backend; kredensial Google dan OpenAI untuk fitur AI."),
]
for x, n, title, body in prereqs:
    rect(s, x, 2.21, 3.42, 3.38, WHITE, radius=True, line=LINE)
    pill(s, n, x + 0.24, 2.47, 0.58, MINT, GREEN, 13)
    txt(s, title, x + 0.24, 3.01, 2.96, 0.67, 19, INK, True)
    txt(s, body, x + 0.24, 3.83, 2.95, 1.16, 13.5, MUTED, valign=MSO_ANCHOR.TOP)
txt(s, "Penting: registrasi menerima URL/ID Google Sheet; pengguna tidak diminta membuat UUID teknis.",
    0.88, 6.13, 11.6, 0.48, 14, GREEN, True)
notes(s, "Untuk audiens bisnis, fokuskan pada tiga prasyarat. Detail port, DSN, dan instalasi server tidak perlu dijelaskan kecuali ada pertanyaan teknis. Konfigurasi AI diperlukan untuk fitur generatif.")

# 05 — role diagram
s = page("Siapa melakukan apa?", "Peran dan akses", 5,
         subtitle="Tindakan pengguna dan ruang data yang boleh diakses diperiksa bersama.")
roles = [
    ("Admin platform", "Akun, atribut, permission, policy"),
    ("Pemilik sumber", "Daftar Sheet, konteks, kepemilikan"),
    ("Data steward", "Klasifikasi, master, taxonomy, kualitas"),
    ("Approver", "Tinjau konfigurasi dan batch"),
    ("Analis / viewer", "Cari data, tabel, dan chart"),
]
for i, (title, body) in enumerate(roles):
    y = 2.03 + i * 0.77
    rect(s, 0.83, y, 3.0, 0.62, DEEP if i < 4 else GREEN, radius=True)
    txt(s, title, 1.05, y + 0.04, 2.58, 0.48, 14, WHITE, True)
    arrow(s, 4.05, y + 0.17, TEAL, 0.22)
    txt(s, body, 4.47, y + 0.03, 4.07, 0.52, 13.5, INK)
rect(s, 9.18, 2.3, 3.2, 3.65, WHITE, radius=True, line=LINE)
txt(s, "Keputusan akses", 9.42, 2.56, 2.7, 0.47, 18, GREEN, True)
for i, label in enumerate(["Role: aksi yang boleh", "Assignment: unit / wilayah", "Policy: sumber / produk", "Waktu berlaku & scope data"]):
    pill(s, label, 9.43, 3.16 + i * 0.59, 2.7, MINT if i < 3 else PALE_AMBER,
         GREEN if i < 3 else AMBER, 10.5)
notes(s, "Jelaskan bahwa role saja tidak otomatis membuka semua data. Admin membuat akun dan policy; pemilik sumber menyiapkan konteks; steward dan approver menjaga kualitas; pengguna analitik melihat data sesuai akses efektif.")

# 06 — process map
s = page("Peta proses: dari Sheet ke keputusan", "Alur utama", 6)
groups = [
    (0.74, "1 · Siapkan", "Akun & akses\nDaftarkan Sheet\nProfiling tab", NAVY),
    (3.88, "2 · Maknai", "Klasifikasi tab\nMaster & taxonomy\nKonfigurasi ETL", GREEN),
    (7.02, "3 · Percayai", "Validasi & review\nPreview import\nApprove & apply", TEAL),
    (10.16, "4 · Gunakan", "Pantau kualitas\nTanya data\nTabel & chart", NAVY),
]
for i, (x, title, body, color) in enumerate(groups):
    rect(s, x, 2.0, 2.45, 3.18, WHITE, radius=True, line=LINE)
    rect(s, x, 2.0, 2.45, 0.7, color, radius=True)
    txt(s, title, x + 0.14, 2.1, 2.17, 0.48, 15, WHITE, True, align=PP_ALIGN.CENTER)
    txt(s, body, x + 0.23, 2.97, 2.0, 1.58, 15, INK, align=PP_ALIGN.CENTER, spacing=1.45)
    if i < 3:
        arrow(s, x + 2.66, 3.43, TEAL, 0.3)
rect(s, 1.0, 5.72, 11.33, 0.7, PALE_AMBER, radius=True)
txt(s, "Gerbang penting: data baru digunakan setelah klasifikasi, validasi, dan keputusan akses selesai.",
    1.22, 5.81, 10.9, 0.51, 15, INK, True, align=PP_ALIGN.CENTER)
notes(s, "Ini slide navigasi utama. Uraikan empat bab besar; tahap master dan taxonomy dipakai bila dataset memang memerlukan rujukan atau kategori baku.")

# 07 — classification decision tree
s = page("Satu spreadsheet bisa berisi dua jenis data", "Klasifikasi tab", 7,
         subtitle="Keputusan dibuat per tab setelah discovery dan profiling, bukan otomatis untuk seluruh file.")
rect(s, 5.05, 2.0, 3.23, 0.72, NAVY, radius=True)
txt(s, "Google Sheet → tab", 5.28, 2.1, 2.75, 0.47, 17, WHITE, True, align=PP_ALIGN.CENTER)
rule(s, 6.67, 2.72, 6.67, 3.08, TEAL, 2)
rect(s, 5.2, 3.08, 2.94, 0.69, MINT, radius=True)
txt(s, "Apakah data ini rujukan?", 5.34, 3.17, 2.67, 0.48, 15, GREEN, True, align=PP_ALIGN.CENTER)
rule(s, 5.2, 3.43, 3.3, 3.92, TEAL, 1.8)
rule(s, 8.14, 3.43, 10.05, 3.92, TEAL, 1.8)
card(s, 0.9, 3.91, 4.55, 1.84, "MASTER · data rujukan",
     "Identitas stabil; contoh: cabang, produk, karyawan. Cari master yang sudah ada dan cocokkan business key.",
     accent=GREEN, title_size=18, body_size=12)
card(s, 7.88, 3.91, 4.55, 1.84, "NON-MASTER · fakta dinamis",
     "Kejadian atau pengukuran; contoh: penjualan dan hasil penimbangan. Relasikan kolom ke master bila perlu.",
     accent=AMBER, title_size=18, body_size=12)
pill(s, "KONFIRMASI PENGGUNA", 5.29, 6.05, 2.76, PALE_AMBER, AMBER, 10)
notes(s, "Contoh satu file berisi tab Produk dan tab Penjualan. Tab Produk adalah master, tab Penjualan adalah non-master. Klasifikasi per tab dikonfirmasi pengguna; aplikasi tidak menganggap semua unggahan sebagai tabel baru yang langsung siap pakai.")

# 08 — master / taxonomy / dynamic relationships
s = page("Bagaimana data saling terhubung?", "Master, taxonomy, dan waktu", 8)
card(s, 0.78, 2.05, 3.3, 1.48, "Master Cabang",
     "Kode cabang → identitas rujukan stabil", accent=GREEN, title_size=17, body_size=12)
card(s, 0.78, 4.18, 3.3, 1.48, "Taxonomy Kategori",
     "Istilah, sinonim, hierarki, versi", accent=TEAL, title_size=17, body_size=12)
rect(s, 5.0, 2.83, 3.35, 2.08, NAVY, radius=True)
txt(s, "Fakta penjualan", 5.26, 3.04, 2.84, 0.58, 21, WHITE, True, align=PP_ALIGN.CENTER)
txt(s, "tanggal  ·  cabang  ·  produk  ·  nilai", 5.31, 3.76, 2.75, 0.65, 13, MINT,
    align=PP_ALIGN.CENTER)
rule(s, 4.08, 2.8, 4.97, 3.55, GREEN, 2)
rule(s, 4.08, 4.93, 4.97, 4.22, TEAL, 2)
arrow(s, 8.7, 3.69, TEAL, 0.33)
card(s, 9.28, 2.74, 3.24, 2.19, "Produk semantik",
     "Dimensi: cabang & waktu\nMetrik: total penjualan\nData siap ditanya & divisualkan", accent=AMBER,
     title_size=17, body_size=13)
txt(s, "Aturan masa berlaku menjaga riwayat saat label, struktur, atau parameter berubah.",
    0.86, 6.19, 11.86, 0.45, 14, GREEN, True, align=PP_ALIGN.CENTER)
notes(s, "Master memberi identitas yang stabil. Taxonomy menyatukan istilah. Fakta menyimpan kejadian pada waktu tertentu. Jangan menimpa histori transaksi hanya karena label master terbaru berubah; beberapa referensi memakai masa berlaku.")

# 09 — import loop
s = page("Data melewati pemeriksaan sebelum dipakai", "Konfigurasi dan import", 9)
steps = [
    (0.7, "01", "Profiling", "Baca struktur & kualitas"),
    (3.25, "02", "Konfigurasi", "Mapping, aturan, metrik"),
    (5.8, "03", "Preview", "Insert, update, konflik"),
    (8.35, "04", "Persetujuan", "Reviewer berbeda"),
    (10.9, "05", "Apply", "Data tepercaya"),
]
for i, (x, n, title, body) in enumerate(steps):
    rect(s, x, 2.31, 1.78, 2.13, WHITE, radius=True, line=LINE)
    oval(s, x + 0.62, 2.51, 0.53, 0.53, MINT)
    txt(s, n, x + 0.68, 2.6, 0.41, 0.3, 11, GREEN, True, align=PP_ALIGN.CENTER)
    txt(s, title, x + 0.14, 3.14, 1.5, 0.39, 14.5, INK, True, align=PP_ALIGN.CENTER)
    txt(s, body, x + 0.14, 3.66, 1.49, 0.54, 11.3, MUTED, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow(s, x + 1.98, 3.19, TEAL, 0.22)
rect(s, 2.56, 5.14, 8.22, 0.99, PALE_AMBER, radius=True)
txt(s, "Jika ambigu / gagal: jawab pertanyaan, koreksi sumber, lalu validasi dan preview ulang.",
    2.82, 5.29, 7.68, 0.67, 15, INK, True, align=PP_ALIGN.CENTER)
rule(s, 6.67, 5.14, 6.67, 4.54, AMBER, 1.8)
notes(s, "Soroti preview sebagai titik keputusan. Sistem menahan apply ketika ada referensi ambigu, kualitas gagal, atau approval belum sesuai. Hasil koreksi dapat ditelusuri ke nilai asal dan alasan.")

# 10 — architecture / guardrails
s = page("Lapisan aplikasi menjaga aliran data", "Fungsi utama", 10)
layers = [
    (0.8, "SUMBER", "Google Sheets", "Discovery & snapshot", NAVY),
    (3.92, "PENGOLAHAN", "ETL + review", "Kualitas, master, taxonomy", GREEN),
    (7.04, "DATA", "Trusted data", "Produk semantik & metrik", TEAL),
    (10.16, "PEMAKAIAN", "Dashboard", "NL2SQL, tabel, chart", NAVY),
]
for i, (x, label, title, body, color) in enumerate(layers):
    rect(s, x, 2.28, 2.39, 2.69, WHITE, radius=True, line=LINE)
    rect(s, x, 2.28, 2.39, 0.51, color, radius=True)
    txt(s, label, x + 0.12, 2.38, 2.15, 0.31, 10, WHITE, True, align=PP_ALIGN.CENTER)
    txt(s, title, x + 0.19, 3.05, 2.03, 0.64, 18, INK, True, align=PP_ALIGN.CENTER)
    txt(s, body, x + 0.22, 3.88, 1.95, 0.66, 12.3, MUTED, align=PP_ALIGN.CENTER)
    if i < 3:
        arrow(s, x + 2.62, 3.48, TEAL, 0.25)
pill(s, "AKSES", 1.2, 5.67, 1.6, MINT, GREEN)
pill(s, "AUDIT", 3.25, 5.67, 1.6, MINT, GREEN)
pill(s, "QUALITY", 5.3, 5.67, 1.6, MINT, GREEN)
pill(s, "JOBS", 7.35, 5.67, 1.6, MINT, GREEN)
pill(s, "AI POLICY", 9.4, 5.67, 1.9, MINT, GREEN)
notes(s, "Lapisan ini membantu audiens memahami fungsi tanpa nama tabel atau endpoint. Jalur data mempunyai kontrol akses, audit, kualitas, pekerjaan terjadwal, dan kebijakan AI di beberapa titik.")

# 11 — NL2SQL
s = page("Dari pertanyaan ke chart", "Analitik bahasa alami", 11,
         subtitle="Pengguna bertanya dalam bahasa biasa; aplikasi memilih data dan bentuk visual yang sesuai.")
rect(s, 0.82, 2.14, 4.05, 1.45, NAVY, radius=True)
txt(s, "“Total penjualan per cabang bulan ini?”", 1.15, 2.35, 3.38, 1.02, 20,
    WHITE, True, align=PP_ALIGN.CENTER)
arrow(s, 5.1, 2.72, TEAL, 0.34)
card(s, 5.68, 2.1, 3.05, 1.55, "Rencana query",
     "Katalog, izin, metrik, dimensi, periode", accent=GREEN, title_size=17, body_size=12)
arrow(s, 8.95, 2.72, TEAL, 0.34)
card(s, 9.46, 2.1, 3.05, 1.55, "Hasil tepercaya",
     "Tabel + rekomendasi chart", accent=AMBER, title_size=17, body_size=12)
rect(s, 1.15, 4.25, 11.03, 1.56, WHITE, radius=True, line=LINE)
chart = [(2.16, 0.5), (3.14, 0.88), (4.12, 0.69), (5.1, 1.15), (6.08, 0.96),
         (7.06, 1.29), (8.04, 0.81)]
for x, height in chart:
    rect(s, x, 5.49 - height, 0.47, height, TEAL, radius=True)
rule(s, 1.73, 5.49, 8.73, 5.49, LINE, 1)
txt(s, "Bar · line · pie · kombinasi", 9.25, 4.68, 2.3, 0.65, 15, GREEN, True)
txt(s, "Jika ambigu, sistem meminta klarifikasi.", 1.2, 6.19, 10.8, 0.43, 13, MUTED)
notes(s, "Contoh pertanyaan yang baik memuat ukuran, dimensi, dan periode. Backend membangun rencana query terstruktur berdasarkan katalog yang boleh diakses; pengguna tidak perlu memasukkan SQL. Chart dapat berupa bar, line, pie, atau kombinasi bila bentuk data mendukung.")

# 12 — Operations and help
s = page("Operasional dan bantuan tetap terlihat", "Sesudah data aktif", 12)
card(s, 0.79, 2.06, 5.68, 3.51, "Pantau kesehatan data",
     "• Status job dan progres\n• Jadwal, dependency, watermark\n• Kualitas, karantina, reprocess\n• Notifikasi yang perlu ditindaklanjuti",
     accent=GREEN, title_size=21, body_size=15)
card(s, 6.83, 2.06, 5.68, 3.51, "Bantu pengguna memahami aplikasi",
     "• Bantuan pada setiap halaman\n• Panduan lengkap 10 tahap\n• Tanya AI atas knowledge base terkurasi\n• Jawaban menyertakan sumber rujukan",
     accent=TEAL, title_size=21, body_size=15)
txt(s, "Asisten membantu penggunaan aplikasi; data analitik tetap ditanya melalui Dashboard / Chat data.",
    0.9, 6.05, 11.5, 0.47, 13.8, GREEN, True, align=PP_ALIGN.CENTER)
notes(s, "Bedakan dua AI untuk audiens: Dashboard/Chat data menjawab pertanyaan tentang angka dan data. Tanya AI menjawab cara memakai aplikasi dari knowledge base terkurasi. Keduanya mengikuti identitas pengguna.")

# 13 — scenario
s = page("Contoh cerita: laporan penjualan cabang", "Skenario demonstrasi", 13)
scenario = [
    (0.82, "Sumber", "Sheet penjualan\ndibagikan & didaftarkan"),
    (3.95, "Konteks", "Tab NON-MASTER\nCabang → master resmi"),
    (7.08, "Kontrol", "Validasi, preview\nreview, lalu apply"),
    (10.21, "Insight", "Total per cabang\nchart bar / line"),
]
for i, (x, title, body) in enumerate(scenario):
    rect(s, x, 2.32, 2.32, 2.13, WHITE, radius=True, line=LINE)
    oval(s, x + 0.85, 2.52, 0.61, 0.61, MINT)
    txt(s, str(i + 1), x + 0.96, 2.63, 0.38, 0.35, 14, GREEN, True, align=PP_ALIGN.CENTER)
    txt(s, title, x + 0.2, 3.22, 1.93, 0.4, 16, INK, True, align=PP_ALIGN.CENTER)
    txt(s, body, x + 0.17, 3.75, 2.0, 0.57, 11.5, MUTED, align=PP_ALIGN.CENTER)
    if i < 3:
        arrow(s, x + 2.53, 3.21, TEAL, 0.24)
rect(s, 1.15, 5.28, 11.02, 0.84, NAVY, radius=True)
txt(s, "Nilai tambah: hasil dapat ditelusuri ke sumber, aturan, waktu, dan persetujuan yang digunakan.",
    1.38, 5.41, 10.55, 0.59, 16, WHITE, True, align=PP_ALIGN.CENTER)
notes(s, "Gunakan satu contoh sederhana agar alur utuh terasa konkret. Bila demonstrasi langsung, siapkan Sheet contoh, akun pemilik dan approver berbeda, serta satu pertanyaan dashboard.")

# 14 — boundaries
s = page("Status dan batas yang perlu disampaikan", "Kesiapan penerapan", 14)
card(s, 0.83, 2.03, 5.65, 3.7, "Tersedia di aplikasi",
     "• Workflow sumber, master, taxonomy, ETL, import, dan dashboard\n• Login, role, policy akses sumber, audit, dan AI helper\n• Dokumentasi API, panduan pengguna, dan pengujian otomatis",
     accent=GREEN, title_size=20, body_size=14)
card(s, 6.85, 2.03, 5.65, 3.7, "Perlu verifikasi saat rollout",
     "• Kredensial dan kuota Google / OpenAI pada lingkungan tujuan\n• Review policy untuk sumber lama dan kontrol akses menyeluruh\n• Uji end-to-end data nyata, performa, backup, dan keamanan produksi",
     accent=AMBER, title_size=20, body_size=14)
txt(s, "Keberhasilan build dan test lokal bukan pengganti uji integrasi serta penerimaan data di production.",
    1.18, 6.05, 11.0, 0.53, 14, RED, True, align=PP_ALIGN.CENTER)
notes(s, "Sampaikan batas secara singkat dan jujur. Kontrol BE16 untuk sumber ber-metadata tersedia, tetapi default-deny seluruh resource dan backfill semua sumber lama belum selesai menurut dokumentasi status. Integrasi Google/OpenAI production perlu diverifikasi terpisah.")

# 15 — filled taxonomy example
s = page("Contoh isian: taxonomy Jenis Biaya", "Contoh form · 1/4", 15,
         subtitle="Contoh fiktif; kode dan label disesuaikan dengan istilah resmi organisasi.")
rect(s, 0.82, 2.1, 3.48, 2.56, NAVY, radius=True)
txt(s, "BUAT TAXONOMY", 1.06, 2.36, 2.92, 0.35, 11, MINT, True)
txt(s, "Kode\njenis_biaya", 1.06, 2.97, 2.86, 0.75, 17, WHITE, True)
txt(s, "Nama\nJenis Biaya", 1.06, 3.89, 2.86, 0.62, 16, WHITE, True)
arrow(s, 4.58, 3.17, TEAL, 0.31)
card(s, 5.13, 2.06, 3.5, 2.58, "Term 1 · operasional",
     "Label: Biaya Operasional\nParent: Root\nAlias: OPEX, biaya rutin", accent=GREEN,
     title_size=17, body_size=14)
card(s, 8.93, 2.06, 3.5, 2.58, "Term 2 · investasi",
     "Label: Biaya Investasi\nParent: Root\nAlias: CAPEX", accent=AMBER,
     title_size=17, body_size=14)
rect(s, 1.03, 5.38, 11.22, 0.76, MINT, radius=True)
txt(s, "Setujui taxonomy  →  pilih kolom “Jenis Biaya”  →  simpan & setujui binding.",
    1.29, 5.49, 10.68, 0.5, 16, GREEN, True, align=PP_ALIGN.CENTER)
notes(s, "Tunjukkan dua lapis pengisian: taxonomy hanya kode dan nama; kemudian setiap term mempunyai kode, label, parent opsional, dan alias opsional. ID internal dibuat sistem. Contoh bukan data yang sudah tersedia pada tenant. Sesudah approval taxonomy, binding kolom juga perlu review.")

# 16 — attribute examples
s = page("Contoh isian: atribut organisasi", "Contoh form · 2/4", 16,
         subtitle="Admin membuat jenis, kode, dan label; pilihan ini kemudian muncul pada form sumber dan akses.")
attributes = [
    (0.83, 2.12, "DEPARTMENT", "keuangan", "Departemen Keuangan", GREEN),
    (4.93, 2.12, "BUSINESS_DOMAIN", "penjualan", "Penjualan", TEAL),
    (9.03, 2.12, "JURISDICTION", "jatim", "Jawa Timur", AMBER),
    (2.87, 4.29, "PURPOSE", "analitik_manajemen", "Analitik Manajemen", GREEN),
    (6.98, 4.29, "CLEARANCE", "internal", "Internal", TEAL),
]
for x, y, kind, code, label, accent in attributes:
    rect(s, x, y, 3.47, 1.75, WHITE, radius=True, line=LINE)
    rect(s, x + 0.18, y + 0.2, 0.07, 1.35, accent, radius=True)
    txt(s, kind, x + 0.39, y + 0.17, 2.83, 0.33, 10, accent, True)
    txt(s, code, x + 0.39, y + 0.58, 2.81, 0.43, 16, INK, True)
    txt(s, label, x + 0.39, y + 1.1, 2.78, 0.39, 12.5, MUTED)
notes(s, "Jelaskan perbedaan: departemen adalah unit, domain adalah proses bisnis, yurisdiksi adalah cakupan wilayah, purpose adalah tujuan penggunaan, clearance adalah atribut kelayakan akses pengguna. Kode boleh dirancang organisasi. Parent opsional untuk hierarki sejenis, tetapi tidak otomatis mewariskan izin.")

# 17 — source metadata example
s = page("Contoh isian: sumber Google Sheet", "Contoh form · 3/4", 17,
         subtitle="Satu contoh pendaftaran; dropdown hanya menampilkan pilihan yang tersedia dan diizinkan.")
card(s, 0.81, 2.04, 5.57, 3.83, "Identitas sumber",
     "Nama: Penjualan Cabang Jawa Timur\nURL: tautan Google Sheet asli\nDeskripsi: transaksi harian per cabang\nJadwal: Tanpa jadwal (awal)",
     accent=GREEN, title_size=20, body_size=16)
card(s, 6.82, 2.04, 5.67, 3.83, "Konteks & tanggung jawab",
     "Unit: Departemen Keuangan\nDomain: Penjualan · Yurisdiksi: Jawa Timur\nPurpose: Analitik Manajemen\nOwner & steward: pilih akun yang tepat\nSensitivitas: LOW jika sesuai klasifikasi",
     accent=TEAL, title_size=20, body_size=14.5)
txt(s, "Kode sumber dan UUID dibuat sistem; spreadsheet harus dibagikan kepada service account sebagai Viewer.",
    0.94, 6.17, 11.9, 0.45, 13.5, GREEN, True, align=PP_ALIGN.CENTER)
notes(s, "Saat demonstrasi, gunakan URL Google Sheet asli yang memang dibagikan kepada service account. Contoh atribut tidak otomatis tersedia; admin perlu membuat dan meng-assign atribut. Pemilik dan steward dipilih dari akun tenant yang tersedia. Sensitivitas LOW hanya contoh dan harus mengikuti isi data nyata.")

# 18 — master, dynamic data, and analysis example
s = page("Contoh isian: master hingga dashboard", "Contoh form · 4/4", 18)
card(s, 0.8, 2.0, 3.46, 2.76, "MASTER · Cabang",
     "Kode registry: cabang\nBusiness key: kode_cabang\nLabel: nama_cabang\nContoh nilai: CBG01 / Malang",
     accent=GREEN, title_size=18, body_size=13.5)
arrow(s, 4.49, 3.18, TEAL, 0.28)
card(s, 5.0, 2.0, 3.46, 2.76, "NON-MASTER · Penjualan",
     "tanggal · kode_cabang\nnomor_transaksi\nnilai_penjualan\nContoh: 2026-10-01 / CBG01",
     accent=TEAL, title_size=18, body_size=13.5)
arrow(s, 8.69, 3.18, TEAL, 0.28)
card(s, 9.19, 2.0, 3.35, 2.76, "ANALITIK",
     "Dimensi: tanggal, cabang\nMetrik: total_penjualan\nTabel / chart sesuai izin",
     accent=AMBER, title_size=18, body_size=13.5)
rect(s, 1.21, 5.38, 10.89, 0.83, NAVY, radius=True)
txt(s, "“Tampilkan total penjualan per cabang bulan ini.”", 1.44, 5.53, 10.38, 0.5,
    19, WHITE, True, align=PP_ALIGN.CENTER)
notes(s, "Business key cabang harus kode stabil, bukan nama yang bisa berubah. Penjualan adalah fakta dinamis dan dapat merujuk master Cabang melalui binding approved. Nama field contoh harus diganti sesuai header asli. Dashboard tetap mengikuti akses pengguna dan status produk.")

# 19 — unit access and reviewers
s = page("Akses multi-unit dan approver", "Tata kelola · keputusan", 19,
         subtitle="Daftar unit untuk melihat data; daftar reviewer untuk keputusan pada sumber.")
card(s, 0.72, 2.07, 3.55, 3.68, "Manajer Penjualan",
     "Unit yang dipilih admin:\n• Penjualan Malang\n• Penjualan Surabaya\n\nInduk organisasi tidak otomatis membuka unit bawahan.",
     accent=GREEN, body_size=14.5)
arrow(s, 4.43, 3.46, TEAL, 0.27)
card(s, 4.93, 2.07, 3.55, 3.68, "Sumber Penjualan Malang",
     "Metadata unit pemilik + domain + yurisdiksi\nPolicy aktif menentukan aksi baca/query\nPenunjukan reviewer tidak membuka data.",
     accent=TEAL, body_size=14.5)
arrow(s, 8.62, 3.46, TEAL, 0.27)
card(s, 9.12, 2.07, 3.49, 3.68, "Reviewer ditunjuk",
     "Metadata: Manajer A\nKonfigurasi: Manajer A\nBatch import: Manajer B\n\nReviewer berbeda dari pengaju.",
     accent=AMBER, body_size=14.5)
txt(s, "Akses = unit eksplisit + policy  |  Approval = reviewer aktif + jenis keputusan + pemisahan tugas",
    0.84, 6.12, 11.8, 0.48, 15, GREEN, True, align=PP_ALIGN.CENTER)
notes(s, "Contoh fiktif. Admin dapat memberikan beberapa unit sekaligus lewat Administrasi. Untuk setiap sumber, admin menunjuk reviewer metadata, konfigurasi, dan batch import secara terpisah. Penunjukan tidak menggantikan policy akses; reviewer harus aktif dan tidak boleh menyetujui pekerjaannya sendiri. Sumber lama yang belum dikonfigurasi masih memakai pemeriksaan role lama sampai daftar reviewer disimpan.")

# 20 — release approvals
s = page("Persetujuan sebelum data tayang", "Tata kelola · gate rilis", 20,
         subtitle="Contoh sumber gabungan: IT dan semua unit terkait menyetujui versi konfigurasi yang sama.")
rect(s, 0.94, 2.18, 11.48, 0.61, NAVY, radius=True)
txt(s, "Konfigurasi ETL versi 2 approved dan siap diperiksa", 1.23, 2.30, 10.86, 0.37,
    17, WHITE, True, align=PP_ALIGN.CENTER)
card(s, 0.8, 2.94, 3.52, 2.18, "IT · teknis",
     "Skema, DQ, keamanan,\nhasil dry-run\nPemeriksa: approver.it",
     accent=GREEN, body_size=14)
card(s, 4.9, 2.94, 3.52, 2.18, "Penjualan Malang",
     "Definisi bisnis, metrik,\nlingkup unit\nPemeriksa: manajer.penjualan",
     accent=TEAL, body_size=14)
card(s, 9.0, 2.94, 3.52, 2.18, "Keuangan Surabaya",
     "Kesesuaian data lintas unit\ndan aturan pemakaian\nPemeriksa: manajer.keuangan",
     accent=AMBER, body_size=14)
txt(s, "+", 4.46, 3.71, 0.34, 0.5, 23, GREEN, True, align=PP_ALIGN.CENTER)
txt(s, "+", 8.56, 3.71, 0.34, 0.5, 23, GREEN, True, align=PP_ALIGN.CENTER)
rect(s, 1.56, 5.65, 10.18, 0.78, MINT, radius=True)
txt(s, "Semua setuju pada revisi yang sama  →  deploy  →  data tersedia sesuai policy akses",
    1.8, 5.83, 9.72, 0.42, 15.4, GREEN, True, align=PP_ALIGN.CENTER)
notes(s, "Admin memilih akun IT dan unit terkait dari daftar. Approver unit wajib memiliki assignment aktif, tetapi penunjukan approval tidak memberi hak query. Jika revisi konfigurasi atau aturan rilis berubah, keputusan lama tidak berlaku. Gate diperiksa pada antrean dan worker deploy; batch import berikutnya tetap mengikuti review batch tersendiri.")

# 21 — close/source
s = page("Pesan utama", "Penutup", 21, dark=True)
txt(s, "Google Sheet tetap menjadi titik awal.\nKepercayaan dibangun melalui proses.",
    0.85, 2.05, 10.65, 1.61, 30, WHITE, True, valign=MSO_ANCHOR.TOP)
txt(s, "Daftarkan  →  Maknai  →  Validasi  →  Setujui  →  Gunakan", 0.91, 4.26, 11.5, 0.8,
    20, MINT, True)
txt(s, "Sumber: USER_GUIDE_END_TO_END.md · ACCESS_JURISDICTION_BE16.md ·\n"
       "MASTER_DATA_DAN_VALIDASI_IMPORT.md · CONTOH_ISIAN_APLIKASI.md · AI_USER_HELP.md",
    0.91, 5.68, 11.2, 0.63, 10, MINT)
notes(s, "Akhiri dengan lima kata kerja. Saat diskusi lanjut, arahkan audiens ke panduan pengguna dan status akses BE16; gunakan deck sebagai pengantar, bukan spesifikasi teknis lengkap.")

prs.core_properties.title = "Google Sheet AI — Gambaran bisnis dan proses aplikasi"
prs.core_properties.subject = "Latar belakang, tujuan, syarat, fungsi, proses, dan diagram aplikasi"
prs.core_properties.author = "Google Sheet AI Project"
prs.core_properties.keywords = "Google Sheets, ETL, master data, taxonomy, NL2SQL, dashboard"
prs.save(OUTPUT)
print(f"Wrote {OUTPUT} ({len(prs.slides)} slides)")
