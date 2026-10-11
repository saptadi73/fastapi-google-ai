"""Create five fictional, typed Excel examples, CSVs and app guidance offline."""

import argparse
import calendar
import csv
import hashlib
import io
import json
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_STORED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUTPUT = ROOT / "docs/examples/analytics"
FIXED_TIME = "2026-03-03T09:00:00+07:00"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


@dataclass
class Formula:
    expression: str
    value: object


def value(item):
    return item.value if isinstance(item, Formula) else item


@dataclass
class Sheet:
    name: str
    grain: str
    rows: list[dict]
    key: str
    role: str = "NON_MASTER"
    metrics: list[tuple] = field(default_factory=list)
    scoped_metrics: list[dict] = field(default_factory=list)
    types: dict = field(default_factory=dict)
    notes: str = ""

    @property
    def headers(self):
        return list(self.rows[0])

    def kind(self, name):
        if name in self.types:
            return self.types[name]
        if name in (
            "period_start",
            "period_end",
            "due_date",
            "followup_start",
            "followup_end",
        ) or name.endswith("_date"):
            return "date"
        if name.endswith("_at"):
            return "timestamptz"
        values = [value(row[name]) for row in self.rows if value(row[name]) is not None]
        if values and all(isinstance(v, bool) for v in values):
            return "boolean"
        if values and all(isinstance(v, int) and not isinstance(v, bool) for v in values):
            return "bigint"
        if values and all(isinstance(v, (float, int)) and not isinstance(v, bool) for v in values):
            return "numeric"
        return "text"

    def config(self, scenario):
        columns = []
        for header in self.headers:
            kind = self.kind(header)
            column = dict(
                source_column=header,
                target_column=header,
                target_type=kind,
                nullable=any(value(row[header]) is None for row in self.rows),
                is_business_key=header == self.key,
                pii_classification="NONE",
                transformation_codes=["parse_date_id"] if kind == "date" else [],
            )
            if kind == "date":
                column.update(date_input_format="TEXT", date_format="%Y-%m-%d")
            columns.append(column)
        metrics = [
            dict(
                code=f"m_{column}_{aggregation}",
                column=column,
                aggregation=aggregation,
                label=label,
                description=description,
                unit=unit,
            )
            for column, aggregation, label, unit, description in self.metrics
        ]
        metrics.extend(self.scoped_metrics)
        dimensions = [
            h
            for h in self.headers
            if self.kind(h) in ("text", "date")
            and h
            not in ("finding", "cause_hypothesis", "action_description", "evidence_ref", "goal_description")
        ][:12]
        quality = []
        for header in self.headers:
            if header.endswith("_score"):
                quality.extend(
                    [dict(column=header, rule="min", value=1), dict(column=header, rule="max", value=5)]
                )
            elif header.startswith("weight_") or header in (
                "response_breach_ratio",
                "csat_ratio",
                "response_rate",
                "attribution_coverage_ratio",
            ):
                quality.extend(
                    [dict(column=header, rule="min", value=0), dict(column=header, rule="max", value=1)]
                )
            elif header.endswith("_count") or header.endswith("_total") or header.endswith("_idr"):
                quality.append(dict(column=header, rule="min", value=0))
        return dict(
            schema_version="1.0",
            dataset_business_name=f"CONTOH FIKTIF {scenario} / {self.name}",
            dataset_description=self.grain + ". " + self.notes,
            grain=self.grain,
            target_table=self.name.lower()[:30],
            target_schema="trusted",
            load_strategy="UPSERT",
            columns=columns,
            data_quality_rules=quality,
            semantic=dict(
                code=f"demo_{scenario}_{self.name}".lower()[:63],
                dimensions=dimensions,
                metrics=metrics,
                allowed_roles=[
                    "PLATFORM_ADMIN",
                    "DATA_STEWARD",
                    "SOURCE_OWNER",
                    "TECHNICAL_APPROVER",
                    "ANALYST",
                    "VIEWER",
                ],
            ),
            unresolved_questions=[
                "Data sintetis: verifikasi grain, key, target, satuan, referensi master, DQ, metrik dan akses sebelum approval."
            ],
            overall_confidence=1,
        )


def period(month):
    return date(2026, month, 1).isoformat(), date(
        2026, month, calendar.monthrange(2026, month)[1]
    ).isoformat()


def stamp(month, day=28, hour=9):
    return f"2026-{month:02d}-{day:02d}T{hour:02d}:00:00+07:00"


def ratio(numerator, denominator):
    return round(numerator / denominator, 6) if denominator else None


def indicator_master():
    definitions = [
        (
            "REVENUE",
            "Pendapatan tercatat",
            "IDR",
            "HIGHER_IS_BETTER",
            "SUM",
            "Jumlah pendapatan tercatat pada periode.",
        ),
        (
            "RESOLVED",
            "Kasus selesai",
            "CASE",
            "HIGHER_IS_BETTER",
            "SUM",
            "Jumlah kasus unik selesai pada periode.",
        ),
        (
            "RESPONSE_HOURS",
            "Jam respons awal",
            "HOUR",
            "LOWER_IS_BETTER",
            "AVG",
            "Rata-rata jam respons dari sampel yang dicatat.",
        ),
        (
            "SATISFACTION",
            "Skor kepuasan",
            "SCORE_1_5",
            "HIGHER_IS_BETTER",
            "AVG",
            "Rating sampel 1 sampai 5, bukan CSAT/NPS.",
        ),
        (
            "SLA_BREACH",
            "Rasio pelanggaran SLA",
            "RATIO",
            "LOWER_IS_BETTER",
            "RECOMPUTE_RATIO",
            "Keluhan melewati SLA / keluhan cohort periode.",
        ),
        (
            "COST_QUALIFIED_LEAD",
            "Biaya per qualified lead",
            "IDR_PER_LEAD",
            "LOWER_IS_BETTER",
            "RECOMPUTE_RATIO",
            "Spend / qualified leads yang teratribusi dan terdeduplikasi.",
        ),
        (
            "PRODUCT_CSAT",
            "CSAT produk",
            "RATIO",
            "HIGHER_IS_BETTER",
            "RECOMPUTE_RATIO",
            "Skor produk 4/5 dibagi respons valid pada versi survei yang sama.",
        ),
        (
            "SERVICE_SCORE",
            "Skor layanan",
            "SCORE_1_5",
            "HIGHER_IS_BETTER",
            "AVG",
            "Rataan skor layanan dalam metode/sampel yang sebanding.",
        ),
    ]
    fields = [
        "indicator_code",
        "indicator_label",
        "measurement_unit",
        "direction",
        "aggregation",
        "definition",
    ]
    return Sheet(
        "Master_Indikator",
        "satu baris per indikator dengan definisi dan satuan",
        [dict(zip(fields, row)) for row in definitions],
        "indicator_code",
        "MASTER",
    )


def refs():
    return {
        "Master_Unit": Sheet(
            "Master_Unit",
            "satu baris per unit",
            [
                dict(unit_code="UNIT-CS", unit_label="Layanan Pelanggan", is_active=True),
                dict(unit_code="UNIT-MKT", unit_label="Marketing", is_active=True),
                dict(unit_code="UNIT-OPS", unit_label="Operasional", is_active=True),
            ],
            "unit_code",
            "MASTER",
        ),
        "Master_Produk": Sheet(
            "Master_Produk",
            "satu baris per produk",
            [
                dict(
                    product_code="PROD-MILK",
                    product_label="Susu Contoh Plain",
                    product_group="DAIRY",
                    is_active=True,
                ),
                dict(
                    product_code="PROD-FEED",
                    product_label="Pakan Contoh",
                    product_group="FEED",
                    is_active=True,
                ),
            ],
            "product_code",
            "MASTER",
        ),
        "Master_Area": Sheet(
            "Master_Area",
            "satu baris per area",
            [
                dict(area_code="AREA-A", area_label="Wilayah Contoh A", is_active=True),
                dict(area_code="AREA-B", area_label="Wilayah Contoh B", is_active=True),
            ],
            "area_code",
            "MASTER",
        ),
        "Master_Tim": Sheet(
            "Master_Tim",
            "satu baris per tim penanggung jawab",
            [
                dict(owner_team_code="TEAM-CS", team_label="Tim Layanan", unit_code="UNIT-CS"),
                dict(owner_team_code="TEAM-MKT", team_label="Tim Marketing", unit_code="UNIT-MKT"),
                dict(owner_team_code="TEAM-OPS", team_label="Tim Operasional", unit_code="UNIT-OPS"),
            ],
            "owner_team_code",
            "MASTER",
        ),
        "Master_Channel": Sheet(
            "Master_Channel",
            "satu baris per channel",
            [
                dict(channel_code="ATL-RADIO", channel_label="Radio contoh", marketing_type="ATL"),
                dict(channel_code="ATL-OOH", channel_label="Media luar ruang contoh", marketing_type="ATL"),
                dict(
                    channel_code="BTL-EVENT", channel_label="Aktivasi lapangan contoh", marketing_type="BTL"
                ),
                dict(
                    channel_code="BTL-DIRECT", channel_label="Kampanye langsung contoh", marketing_type="BTL"
                ),
            ],
            "channel_code",
            "MASTER",
        ),
        "Master_Competitor": Sheet(
            "Master_Competitor",
            "satu baris per pelaku pasar",
            [
                dict(competitor_code="OWN", competitor_label="Organisasi Kita Contoh", is_own=True),
                dict(competitor_code="COMP-A", competitor_label="Pesaing Fiktif A", is_own=False),
                dict(competitor_code="COMP-B", competitor_label="Pesaing Fiktif B", is_own=False),
            ],
            "competitor_code",
            "MASTER",
        ),
    }


def action(
    code,
    sheet,
    row_id,
    unit,
    team,
    indicator,
    baseline,
    target,
    unit_measure,
    finding,
    hypothesis,
    description,
):
    return dict(
        action_id=code,
        source_sheet=sheet,
        source_record_id=row_id,
        unit_code=unit,
        indicator_code=indicator,
        period_start="2026-01-01",
        period_end="2026-01-31",
        baseline_value=baseline,
        target_value=target,
        measurement_unit=unit_measure,
        finding=finding,
        cause_hypothesis=hypothesis,
        cause_status="UNVERIFIED",
        evidence_ref=f"SIMULASI::{sheet}::{row_id}",
        action_description=description,
        owner_team_code=team,
        due_date="2026-02-15",
        action_status="PLANNED",
        followup_start="2026-02-01",
        followup_end="2026-02-28",
        followup_value=None,
        verification_status="NOT_EVALUATED",
        collected_at="2026-02-02T09:00:00+07:00",
    )


def action_sheet(rows):
    return Sheet(
        "Action_Plan",
        "satu baris per tindakan yang merujuk record temuan",
        rows,
        "action_id",
        types={"baseline_value": "numeric", "target_value": "numeric", "followup_value": "numeric"},
        notes="Baseline dibekukan sesuai periode bukti; hipotesis belum terverifikasi. Nilai follow-up kosong berarti belum dievaluasi, bukan nol.",
    )


def row_formula(headers, row_number, output, expression, cached):
    positions = {h: f"{get_column_letter(i + 1)}{row_number}" for i, h in enumerate(headers)}
    return Formula("=" + expression.format(**positions), cached)


def aggregate_formula(source, target, row_number, column, operation, cached, matches):
    """Use whole-column formulas so added source rows remain inside the aggregate."""
    source_columns = {
        h: f"'{source.name}'!{get_column_letter(i + 1)}:{get_column_letter(i + 1)}"
        for i, h in enumerate(source.headers)
    }
    target_cells = {h: f"{get_column_letter(i + 1)}{row_number}" for i, h in enumerate(target)}
    arguments = []
    for source_column, target_column in matches:
        arguments.extend([source_columns[source_column], target_cells[target_column]])
    if operation == "COUNTIFS":
        expression = operation + "(" + ",".join(arguments) + ")"
    else:
        expression = operation + "(" + ",".join([source_columns[column], *arguments]) + ")"
    return Formula("=" + expression, cached)


def complaints():
    rows = []
    for month in (1, 2):
        for i in range(16):
            start = datetime(2026, month, i + 2, 9, tzinfo=timezone(timedelta(hours=7)))
            hours = [4, 12, 30, 40][i % 4] if month == 1 else [3, 6, 12, 28][i % 4]
            first = start + timedelta(hours=hours)
            resolved = None if i % 7 == 0 else first + timedelta(hours=12 + i)
            rows.append(
                dict(
                    complaint_id=f"CLM-{month:02d}-{i + 1:03d}",
                    unit_code="UNIT-CS",
                    product_code="PROD-MILK" if i % 2 == 0 else "PROD-FEED",
                    area_code="AREA-A" if i % 3 else "AREA-B",
                    category_code="SERVICE" if i % 2 == 0 else "QUALITY",
                    priority_code="HIGH" if i % 5 == 0 else "NORMAL",
                    reported_at=start.isoformat(),
                    first_response_at=first.isoformat(),
                    resolved_at=resolved.isoformat() if resolved else None,
                    complaint_status="RESOLVED" if resolved else "OPEN",
                    first_response_hours=hours,
                    resolution_hours=round((resolved - start).total_seconds() / 3600, 2)
                    if resolved
                    else None,
                    response_sla_hours=24,
                    complaint_count=1,
                    response_breach_count=int(hours > 24),
                    resolved_count=int(resolved is not None),
                    period_start=period(month)[0],
                    period_end=period(month)[1],
                    collected_at=stamp(month + 1, 2),
                    data_as_of_at=stamp(month, calendar.monthrange(2026, month)[1], 23),
                )
            )
    detail = Sheet(
        "Keluhan_Data",
        "satu baris per keluhan anonim",
        rows,
        "complaint_id",
        metrics=[
            (
                "complaint_count",
                "sum",
                "Jumlah keluhan",
                "CASE",
                "Tidak menjumlahkan ulang rekap dengan detail.",
            ),
            (
                "response_breach_count",
                "sum",
                "Keluhan melewati SLA respons",
                "CASE",
                "Respons awal lebih dari 24 jam pada contoh.",
            ),
            (
                "first_response_hours",
                "avg",
                "Rata-rata jam respons awal",
                "HOUR",
                "Semua contoh telah memiliki respons awal.",
            ),
        ],
    )
    summary = []
    for month in (1, 2):
        for product in ("PROD-MILK", "PROD-FEED"):
            matching = [
                r
                for r in rows
                if r["reported_at"].startswith(f"2026-{month:02d}") and r["product_code"] == product
            ]
            breaches = sum(r["response_breach_count"] for r in matching)
            summary.append(
                dict(
                    complaint_summary_id=f"CLS-{month:02d}-{product}",
                    unit_code="UNIT-CS",
                    product_code=product,
                    period_start=period(month)[0],
                    period_end=period(month)[1],
                    complaints_total=len(matching),
                    resolved_total=sum(r["resolved_count"] for r in matching),
                    response_breach_total=breaches,
                    response_breach_ratio=ratio(breaches, len(matching)),
                    target_max_breach_ratio=0.1,
                    avg_first_response_hours=round(
                        sum(r["first_response_hours"] for r in matching) / len(matching), 2
                    ),
                    reporting_status="FINAL_SIMULATED",
                    collected_at=stamp(month + 1, 2),
                    data_as_of_at=stamp(month, calendar.monthrange(2026, month)[1], 23),
                )
            )
    for index, row in enumerate(summary, 2):
        for output, source, operation in [
            ("complaints_total", "complaint_count", "SUMIFS"),
            ("resolved_total", "resolved_count", "SUMIFS"),
            ("response_breach_total", "response_breach_count", "SUMIFS"),
            ("avg_first_response_hours", "first_response_hours", "AVERAGEIFS"),
        ]:
            row[output] = aggregate_formula(
                detail,
                row,
                index,
                source,
                operation,
                row[output],
                [
                    ("product_code", "product_code"),
                    ("unit_code", "unit_code"),
                    ("period_start", "period_start"),
                    ("period_end", "period_end"),
                ],
            )
        row["response_breach_ratio"] = row_formula(
            list(row),
            index,
            "response_breach_ratio",
            'IF({complaints_total}>0,{response_breach_total}/{complaints_total},"")',
            row["response_breach_ratio"],
        )
    recap = Sheet(
        "Keluhan_Rekap",
        "satu baris per unit produk bulan",
        summary,
        "complaint_summary_id",
        metrics=[
            (
                "complaints_total",
                "sum",
                "Jumlah keluhan rekap",
                "CASE",
                "Pilih rekap saja atau detail saja, jangan gabungkan keduanya dalam total.",
            )
        ],
        notes="Rekap memakai SUMIFS/AVERAGEIFS dari detail dan formula pembilang/penyebut. Keluhan terselesaikan adalah dari cohort keluhan yang diterima pada periode tersebut, bukan semua penyelesaian periode.",
    )
    actions = action_sheet(
        [
            action(
                "ACT-CLM-01",
                recap.name,
                summary[0]["complaint_summary_id"],
                "UNIT-CS",
                "TEAM-CS",
                "SLA_BREACH",
                value(summary[0]["response_breach_ratio"]),
                0.1,
                "RATIO",
                "Rasio respons lewat SLA melebihi batas contoh 10%.",
                "Antrean atau pembagian shift mungkin berpengaruh; perlu bukti operasional.",
                "Verifikasi antrean dan shift, uji penjadwalan terbatas, lalu ukur SLA pada periode follow-up.",
            )
        ]
    )
    return [detail, recap, actions], ["Master_Unit", "Master_Produk", "Master_Area", "Master_Tim"]


def marketing():
    rows = []
    channels = refs()["Master_Channel"].rows
    for month in (1, 2):
        for i, channel in enumerate(channels):
            for area in ("AREA-A", "AREA-B"):
                attributed = channel["marketing_type"] == "BTL"
                spend = (i + 1) * 500000 + (100000 if area == "AREA-B" else 0)
                leads = 20 + i * 5 + month if attributed else None
                qualified = leads // 2 if attributed else None
                orders = qualified // 3 if attributed else None
                revenue = orders * 250000 if attributed else None
                row = dict(
                    marketing_record_id=f"MKT-{month}-{i}-{area}",
                    campaign_id=f"CMP-{i}-{area}",
                    unit_code="UNIT-MKT",
                    product_code="PROD-MILK",
                    area_code=area,
                    channel_code=channel["channel_code"],
                    marketing_type=channel["marketing_type"],
                    period_start=period(month)[0],
                    period_end=period(month)[1],
                    budget_idr=2500000,
                    spend_idr=spend,
                    reach_count=5000 * (i + 1),
                    reach_method="ESTIMATED" if not attributed else "PLATFORM_REPORTED",
                    leads_count=leads,
                    qualified_leads_count=qualified,
                    orders_count=orders,
                    attributed_revenue_idr=revenue,
                    attribution_model="FIRST_TOUCH_7D_DEDUP" if attributed else "UNATTRIBUTED",
                    attribution_coverage_ratio=1.0 if attributed else None,
                    cpl_idr=ratio(spend, leads) if attributed else None,
                    cost_per_qualified_lead_idr=ratio(spend, qualified) if attributed else None,
                    roas_ratio=ratio(revenue, spend) if attributed else None,
                    evidence_ref=f"SIMULASI-CAMPAIGN-{month}-{i}-{area}",
                    reporting_status="FINAL_SIMULATED",
                    collected_at=FIXED_TIME,
                    data_as_of_at=FIXED_TIME,
                )
                index = len(rows) + 2
                row["cpl_idr"] = row_formula(
                    list(row),
                    index,
                    "cpl_idr",
                    'IF(AND(ISNUMBER({leads_count}),{leads_count}>0),{spend_idr}/{leads_count},"")',
                    row["cpl_idr"],
                )
                row["cost_per_qualified_lead_idr"] = row_formula(
                    list(row),
                    index,
                    "cost_per_qualified_lead_idr",
                    'IF(AND(ISNUMBER({qualified_leads_count}),{qualified_leads_count}>0),{spend_idr}/{qualified_leads_count},"")',
                    row["cost_per_qualified_lead_idr"],
                )
                row["roas_ratio"] = row_formula(
                    list(row),
                    index,
                    "roas_ratio",
                    'IF(AND(ISNUMBER({attributed_revenue_idr}),{spend_idr}>0),{attributed_revenue_idr}/{spend_idr},"")',
                    row["roas_ratio"],
                )
                rows.append(row)
    fact = Sheet(
        "Marketing_Data",
        "satu baris per campaign channel area produk periode",
        rows,
        "marketing_record_id",
        metrics=[
            ("spend_idr", "sum", "Biaya marketing", "IDR", "Belanja aktual pada periode, bukan anggaran."),
            (
                "qualified_leads_count",
                "sum",
                "Leads berkualitas teratribusi",
                "LEAD",
                "Hanya jika deduplikasi dan attribution model konsisten; NULL ATL tidak sama dengan nol.",
            ),
        ],
        types={
            "leads_count": "bigint",
            "qualified_leads_count": "bigint",
            "orders_count": "bigint",
            "attributed_revenue_idr": "numeric",
        },
        notes="ATL/BTL adalah klasifikasi contoh; sepakati dengan bisnis. Reach antar-channel dapat overlap, bukan jangkauan unik gabungan. Rasio total harus dihitung ulang dari pembilang/penyebut, bukan AVG CPL/ROAS.",
    )
    baseline_row = next(
        row for row in rows if row["channel_code"] == "BTL-EVENT" and row["area_code"] == "AREA-A"
    )
    actions = action_sheet(
        [
            action(
                "ACT-MKT-01",
                fact.name,
                baseline_row["marketing_record_id"],
                "UNIT-MKT",
                "TEAM-MKT",
                "COST_QUALIFIED_LEAD",
                value(baseline_row["cost_per_qualified_lead_idr"]),
                80000,
                "IDR_PER_LEAD",
                "Biaya per qualified lead perlu dibandingkan dengan batas biaya contoh.",
                "Kesesuaian audiens/penawaran masih hipotesis, bukan bukti sebab-akibat.",
                "Audit deduplikasi dan attribution, lalu uji audiens kecil dengan anggaran terkendali; ukur qualified lead dan biaya.",
            )
        ]
    )
    return [fact, actions], ["Master_Unit", "Master_Produk", "Master_Area", "Master_Channel", "Master_Tim"]


def satisfaction():
    rows = []
    for month in (1, 2):
        for product in ("PROD-MILK", "PROD-FEED"):
            for i in range(12):
                valid = i != 11
                score = (
                    ([2, 3, 3, 4, 4, 5][i % 6] if month == 1 else [3, 4, 4, 4, 5, 5][i % 6])
                    if valid
                    else None
                )
                rows.append(
                    dict(
                        response_id=f"SUR-{month}-{product}-{i:02d}",
                        respondent_code=f"ANON-{month}-{product}-{i:02d}",
                        unit_code="UNIT-CS",
                        product_code=product,
                        area_code="AREA-A",
                        survey_version="CSAT-5-V1",
                        surveyed_at=stamp(month, i + 10),
                        product_score=score,
                        service_score=4 if valid else None,
                        value_score=3 if valid else None,
                        is_valid_response=valid,
                        excluded_reason=None if valid else "INCOMPLETE",
                        valid_response_count=int(valid),
                        satisfied_count=int(valid and score >= 4),
                        sampling_method="OPT_IN_SIMULATED",
                        period_start=period(month)[0],
                        period_end=period(month)[1],
                        collected_at=FIXED_TIME,
                        data_as_of_at=FIXED_TIME,
                    )
                )
    raw = Sheet(
        "Survey_Data",
        "satu baris per respons anonim produk metode survei",
        rows,
        "response_id",
        metrics=[
            (
                "valid_response_count",
                "sum",
                "Jumlah respons valid",
                "RESPONSE",
                "Respons tidak lengkap tidak masuk penyebut CSAT.",
            ),
            (
                "satisfied_count",
                "sum",
                "Respons puas",
                "RESPONSE",
                "Skor produk 4 atau 5 pada skala 1 sampai 5.",
            ),
        ],
        notes="CSAT produk berdasarkan product_score >=4; bukan NPS. Jangan mencampur versi pertanyaan atau menganggap sampel opt-in mewakili seluruh pelanggan.",
    )
    summary = []
    for month in (1, 2):
        for product in ("PROD-MILK", "PROD-FEED"):
            group = [
                r
                for r in rows
                if r["product_code"] == product and r["surveyed_at"].startswith(f"2026-{month:02d}")
            ]
            valid = [r for r in group if r["is_valid_response"]]
            satisfied = sum(r["satisfied_count"] for r in valid)
            row = dict(
                csat_summary_id=f"CSAT-{month}-{product}",
                unit_code="UNIT-CS",
                product_code=product,
                area_code="AREA-A",
                period_start=period(month)[0],
                period_end=period(month)[1],
                survey_version="CSAT-5-V1",
                invited_count=100,
                received_count=len(group),
                valid_response_count=len(valid),
                excluded_count=len(group) - len(valid),
                satisfied_count=satisfied,
                csat_ratio=ratio(satisfied, len(valid)),
                target_csat_ratio=0.8,
                avg_product_score=round(sum(r["product_score"] for r in valid) / len(valid), 6),
                response_rate=ratio(len(group), 100),
                sample_status="SMALL_OPT_IN_SAMPLE",
                reporting_status="FINAL_SIMULATED",
                collected_at=FIXED_TIME,
                data_as_of_at=FIXED_TIME,
            )
            for output, source, operation in [
                ("received_count", "response_id", "COUNTIFS"),
                ("valid_response_count", "valid_response_count", "SUMIFS"),
                ("satisfied_count", "satisfied_count", "SUMIFS"),
            ]:
                row[output] = aggregate_formula(
                    raw,
                    row,
                    len(summary) + 2,
                    source,
                    operation,
                    row[output],
                    [
                        ("product_code", "product_code"),
                        ("area_code", "area_code"),
                        ("survey_version", "survey_version"),
                        ("period_start", "period_start"),
                        ("period_end", "period_end"),
                    ],
                )
            row["excluded_count"] = row_formula(
                list(row),
                len(summary) + 2,
                "excluded_count",
                "{received_count}-{valid_response_count}",
                row["excluded_count"],
            )
            valid_col = get_column_letter(raw.headers.index("is_valid_response") + 1)
            avg_formula = aggregate_formula(
                raw,
                row,
                len(summary) + 2,
                "product_score",
                "AVERAGEIFS",
                row["avg_product_score"],
                [
                    ("product_code", "product_code"),
                    ("area_code", "area_code"),
                    ("survey_version", "survey_version"),
                    ("period_start", "period_start"),
                    ("period_end", "period_end"),
                ],
            )
            avg_formula.expression = (
                avg_formula.expression[:-1] + f",'{raw.name}'!{valid_col}:{valid_col},TRUE)"
            )
            row["avg_product_score"] = avg_formula
            row["csat_ratio"] = row_formula(
                list(row),
                len(summary) + 2,
                "csat_ratio",
                'IF({valid_response_count}>0,{satisfied_count}/{valid_response_count},"")',
                row["csat_ratio"],
            )
            row["response_rate"] = row_formula(
                list(row),
                len(summary) + 2,
                "response_rate",
                'IF({invited_count}>0,{received_count}/{invited_count},"")',
                row["response_rate"],
            )
            summary.append(row)
    recap = Sheet(
        "CSAT_Rekap",
        "satu baris per produk area versi survei periode",
        summary,
        "csat_summary_id",
        metrics=[
            (
                "valid_response_count",
                "sum",
                "Respons valid rekap",
                "RESPONSE",
                "Jangan menjumlahkan kembali raw dan rekap.",
            ),
            (
                "satisfied_count",
                "sum",
                "Respons puas rekap",
                "RESPONSE",
                "Rasio gabungan perlu jumlah puas / jumlah valid, bukan AVG persentase.",
            ),
        ],
        notes="Sampel kecil dan opt-in: tidak ada margin of error atau klaim representatif otomatis. Rekap memakai respons valid dan formula dari raw; invited_count harus diperbarui dari log undangan yang sah.",
    )
    actions = action_sheet(
        [
            action(
                "ACT-CSAT-01",
                recap.name,
                summary[0]["csat_summary_id"],
                "UNIT-CS",
                "TEAM-CS",
                "PRODUCT_CSAT",
                value(summary[0]["csat_ratio"]),
                0.8,
                "RATIO",
                "CSAT produk Januari di bawah target contoh 80%.",
                "Masalah pengalaman produk masih dugaan; sampel kecil/self-selection dapat memengaruhi hasil.",
                "Perbaiki cakupan sampel dan lakukan wawancara terstruktur anonim; verifikasi tema masalah sebelum uji perbaikan produk.",
            )
        ]
    )
    return [raw, recap, actions], ["Master_Unit", "Master_Produk", "Master_Area", "Master_Tim"]


def competitiveness():
    prices, responses = [], []
    for month in (1, 2):
        for brand_index, brand in enumerate(("OWN", "COMP-A", "COMP-B")):
            volume = (1000, 900, 1000)[brand_index]
            price = (26000, 24000, 28000)[brand_index] + (500 if month == 2 else 0)
            row = dict(
                price_observation_id=f"PRICE-{month}-{brand}",
                competitor_code=brand,
                product_code="PROD-MILK",
                area_code="AREA-A",
                comparison_group="PLAIN_MILK_EQUIVALENT",
                package_volume_ml=volume,
                package_price_idr=price,
                normalized_price_idr_per_l=round(price * 1000 / volume, 6),
                is_promo=False,
                is_in_stock=True,
                outlet_type="RETAIL_SIMULATED",
                observed_at=stamp(month),
                period_start=period(month)[0],
                period_end=period(month)[1],
                evidence_ref=f"SIMULASI-PRICE-{month}-{brand}",
                collected_at=FIXED_TIME,
                data_as_of_at=stamp(month),
            )
            row["normalized_price_idr_per_l"] = row_formula(
                list(row),
                len(prices) + 2,
                "normalized_price_idr_per_l",
                'IF({package_volume_ml}>0,{package_price_idr}*1000/{package_volume_ml},"")',
                row["normalized_price_idr_per_l"],
            )
            prices.append(row)
            for i in range(12):
                product_score = [3, 4, 4, 5][(i + brand_index) % 4]
                service_score = [2, 3, 4, 4][i % 4] if brand == "OWN" else [3, 4, 4, 5][i % 4]
                value_score = [3, 3, 4, 4][i % 4]
                response = dict(
                    competitive_response_id=f"CMP-SUR-{month}-{brand}-{i:02d}",
                    competitor_code=brand,
                    product_code="PROD-MILK",
                    area_code="AREA-A",
                    survey_version="COMPARE-5-V1",
                    surveyed_at=stamp(month, i + 15),
                    product_score=product_score,
                    service_score=service_score,
                    value_score=value_score,
                    weight_product=0.5,
                    weight_service=0.3,
                    weight_value=0.2,
                    weighted_score=round(product_score * 0.5 + service_score * 0.3 + value_score * 0.2, 6),
                    sampling_method="MATCHED_SCRIPT_SIMULATED",
                    period_start=period(month)[0],
                    period_end=period(month)[1],
                    collected_at=FIXED_TIME,
                )
                response["weighted_score"] = row_formula(
                    list(response),
                    len(responses) + 2,
                    "weighted_score",
                    "{product_score}*{weight_product}+{service_score}*{weight_service}+{value_score}*{weight_value}",
                    response["weighted_score"],
                )
                responses.append(response)
    price_sheet = Sheet(
        "Harga_Observasi",
        "satu baris per brand produk ekuivalen outlet area waktu observasi",
        prices,
        "price_observation_id",
        metrics=[
            (
                "normalized_price_idr_per_l",
                "avg",
                "Harga observasi per liter",
                "IDR_PER_L",
                "Hanya dalam comparison_group, area, waktu, promo dan outlet yang sebanding; bukan harga transaksi tertimbang volume.",
            )
        ],
        notes="Harga adalah titik observasi, bukan rata-rata sebulan. COMP-A 900 ml harus dinormalisasi sebelum dibandingkan dengan 1000 ml. product_code menunjukkan kelompok referensi, bukan klaim SKU pesaing identik.",
    )
    survey = Sheet(
        "Competitive_Survey",
        "satu baris per respons brand produk area versi survei",
        responses,
        "competitive_response_id",
        metrics=[
            (
                "weighted_score",
                "avg",
                "Skor gabungan sesuai bobot contoh",
                "SCORE_1_5",
                "Bobot 0.5/0.3/0.2 sama untuk semua brand; rating bukan bukti sebab-akibat.",
            )
        ],
    )
    summaries = []
    for price_row in prices:
        month = int(price_row["period_start"][5:7])
        brand = price_row["competitor_code"]
        group = [
            r
            for r in responses
            if r["competitor_code"] == brand and r["surveyed_at"].startswith(f"2026-{month:02d}")
        ]
        summaries.append(
            dict(
                competitive_summary_id=f"COMP-{month}-{brand}",
                competitor_code=brand,
                product_code="PROD-MILK",
                area_code="AREA-A",
                period_start=period(month)[0],
                period_end=period(month)[1],
                price_observation_id=price_row["price_observation_id"],
                normalized_price_idr_per_l=value(price_row["normalized_price_idr_per_l"]),
                price_observed_at=price_row["observed_at"],
                sample_count=len(group),
                avg_service_score=round(sum(r["service_score"] for r in group) / len(group), 6),
                avg_weighted_score=round(sum(value(r["weighted_score"]) for r in group) / len(group), 6),
                survey_version="COMPARE-5-V1",
                sample_status="SMALL_SAMPLE",
                collected_at=FIXED_TIME,
            )
        )
    for row_index, row in enumerate(summaries, 2):
        row["normalized_price_idr_per_l"] = aggregate_formula(
            price_sheet,
            row,
            row_index,
            "normalized_price_idr_per_l",
            "SUMIFS",
            row["normalized_price_idr_per_l"],
            [("price_observation_id", "price_observation_id")],
        )
        for output, source, operation in [
            ("sample_count", "competitive_response_id", "COUNTIFS"),
            ("avg_service_score", "service_score", "AVERAGEIFS"),
            ("avg_weighted_score", "weighted_score", "AVERAGEIFS"),
        ]:
            row[output] = aggregate_formula(
                survey,
                row,
                row_index,
                source,
                operation,
                row[output],
                [
                    ("competitor_code", "competitor_code"),
                    ("product_code", "product_code"),
                    ("area_code", "area_code"),
                    ("survey_version", "survey_version"),
                    ("period_start", "period_start"),
                    ("period_end", "period_end"),
                ],
            )
    recap = Sheet(
        "Competitive_Rekap",
        "satu baris per brand produk area versi survei periode",
        summaries,
        "competitive_summary_id",
        notes="Baca harga titik observasi dan skor periode sebagai dua bukti berbeda. Tidak ada rank/indeks gabungan harga-layanan otomatis; bobot harus disetujui bisnis.",
    )
    actions = action_sheet(
        [
            action(
                "ACT-COMP-01",
                recap.name,
                summaries[0]["competitive_summary_id"],
                "UNIT-CS",
                "TEAM-CS",
                "SERVICE_SCORE",
                value(summaries[0]["avg_service_score"]),
                4.0,
                "SCORE_1_5",
                "Skor layanan organisasi contoh perlu dibandingkan pada metode yang sama.",
                "Proses layanan mungkin berpengaruh; persepsi brand dan bias sampel belum dieliminasi.",
                "Verifikasi kesetaraan produk/outlet dan tambah sampel; audit proses layanan sebelum uji perbaikan. Jangan mengubah harga hanya dari skor contoh.",
            )
        ]
    )
    return [price_sheet, survey, recap, actions], [
        "Master_Unit",
        "Master_Produk",
        "Master_Area",
        "Master_Competitor",
        "Master_Tim",
    ]


def kpi():
    indicators = indicator_master()
    plans, measurements, results = [], [], []
    targets = [3000000, 30, 24, 4.0]
    for month in (1, 2):
        for index, indicator in enumerate(indicators.rows[:4]):
            unit = "UNIT-MKT" if index == 0 else "UNIT-CS"
            plan_id = f"PLAN-{month}-{indicator['indicator_code']}"
            target = targets[index]
            plans.append(
                dict(
                    plan_id=plan_id,
                    goal_id="GOAL-EXPERIENCE-2026",
                    goal_description="Tingkatkan kinerja bisnis terukur sesuai indikator",
                    unit_code=unit,
                    indicator_code=indicator["indicator_code"],
                    measurement_unit=indicator["measurement_unit"],
                    direction=indicator["direction"],
                    period_start=period(month)[0],
                    period_end=period(month)[1],
                    target_value=target,
                    target_version="TARGET-V1-FICTITIOUS",
                    target_set_at="2025-12-15T09:00:00+07:00",
                    collected_at=FIXED_TIME,
                )
            )
            observed = (
                [700000, 800000, 900000]
                if index == 0
                else [7, 8, 9]
                if index == 1
                else [28, 30, 32]
                if index == 2
                else [3, 4, 3]
            )
            if month == 2:
                observed = [x + 1 if index == 1 else x * (0.8 if index == 2 else 1.1) for x in observed]
            for i, observed_value in enumerate(observed):
                measurements.append(
                    dict(
                        measurement_id=f"MEAS-{month}-{index}-{i}",
                        plan_id=plan_id,
                        unit_code=unit,
                        indicator_code=indicator["indicator_code"],
                        measured_at=stamp(month, i + 12),
                        measured_value=round(observed_value, 6),
                        measurement_unit=indicator["measurement_unit"],
                        evidence_ref=f"SIMULASI-MEAS-{month}-{index}-{i}",
                        method_version="METHOD-V1",
                        collected_at=FIXED_TIME,
                    )
                )
            actual = sum(observed) if indicator["aggregation"] == "SUM" else sum(observed) / len(observed)
            higher = indicator["direction"] == "HIGHER_IS_BETTER"
            row = dict(
                kpi_result_id=f"RESULT-{month}-{indicator['indicator_code']}",
                plan_id=plan_id,
                goal_id="GOAL-EXPERIENCE-2026",
                unit_code=unit,
                indicator_code=indicator["indicator_code"],
                period_start=period(month)[0],
                period_end=period(month)[1],
                measurement_unit=indicator["measurement_unit"],
                direction=indicator["direction"],
                aggregation=indicator["aggregation"],
                target_value=target,
                actual_value=round(actual, 6),
                sample_count=len(observed),
                unfavourable_gap=round(max(0, target - actual if higher else actual - target), 6),
                achievement_ratio=ratio(actual, target) if higher else ratio(target, actual),
                result_status="ON_TARGET"
                if (actual >= target if higher else actual <= target)
                else "BELOW_TARGET",
                reporting_status="FINAL_SIMULATED",
                collected_at=FIXED_TIME,
                data_as_of_at=FIXED_TIME,
            )
            row["unfavourable_gap"] = row_formula(
                list(row),
                len(results) + 2,
                "unfavourable_gap",
                'IF({direction}="HIGHER_IS_BETTER",MAX(0,{target_value}-{actual_value}),MAX(0,{actual_value}-{target_value}))',
                row["unfavourable_gap"],
            )
            row["achievement_ratio"] = row_formula(
                list(row),
                len(results) + 2,
                "achievement_ratio",
                'IF({direction}="HIGHER_IS_BETTER",IF({target_value}>0,{actual_value}/{target_value},""),IF({actual_value}>0,{target_value}/{actual_value},""))',
                row["achievement_ratio"],
            )
            results.append(row)
    plan = Sheet(
        "KPI_Plan",
        "satu baris per unit indikator periode versi target aktif",
        plans,
        "plan_id",
        types={"target_value": "numeric"},
        notes="goal_id bukan business key karena satu goal dapat memiliki banyak indikator/periode. Versi target direview sebelum evaluasi.",
    )
    observation = Sheet(
        "KPI_Measurement",
        "satu baris per bukti pengukuran indikator",
        measurements,
        "measurement_id",
        notes="Nilai berbeda satuan tidak dijumlahkan. Rekap aktual mengikuti SUM/AVG pada Master_Indikator; 3 sampel tidak mewakili populasi tanpa bukti.",
    )
    result = Sheet(
        "KPI_Result",
        "satu baris per plan indikator periode",
        results,
        "kpi_result_id",
        types={"target_value": "numeric", "actual_value": "numeric"},
        notes="Jangan SUM actual/gap/achievement lintas indikator atau AVG rasio tanpa bobot yang sah. Actual rekap perlu dihitung ulang jika pengukuran berubah; rasio arah LOWER berbeda dari HIGHER.",
    )
    for indicator in indicators.rows[:4]:
        scope = [dict(field="indicator_code", operator="eq", value=indicator["indicator_code"])]
        observation.scoped_metrics.append(
            dict(
                code="m_" + indicator["indicator_code"].lower() + "_observed",
                column="measured_value",
                aggregation=indicator["aggregation"].lower(),
                label=indicator["indicator_label"] + " dari bukti",
                unit=indicator["measurement_unit"],
                description="Filter indikator mencegah campuran satuan; pertahankan unit/periode dan metode yang sebanding.",
                filters=scope,
            )
        )
        for column in ("target_value", "actual_value"):
            result.scoped_metrics.append(
                dict(
                    code="m_" + indicator["indicator_code"].lower() + "_" + column,
                    column=column,
                    aggregation="max",
                    label=indicator["indicator_label"] + " " + column,
                    unit=indicator["measurement_unit"],
                    filters=scope,
                    description="MAX hanya untuk menampilkan satu record plan/unit/indikator/periode yang dipilih. Bukan total atau rataan lintas plan/unit/periode.",
                )
            )
    for row_index, row in enumerate(results, 2):
        row["target_value"] = aggregate_formula(
            plan, row, row_index, "target_value", "SUMIFS", row["target_value"], [("plan_id", "plan_id")]
        )
        row["actual_value"] = aggregate_formula(
            observation,
            row,
            row_index,
            "measured_value",
            "SUMIFS" if row["aggregation"] == "SUM" else "AVERAGEIFS",
            row["actual_value"],
            [("plan_id", "plan_id")],
        )
        row["sample_count"] = aggregate_formula(
            observation,
            row,
            row_index,
            "measurement_id",
            "COUNTIFS",
            row["sample_count"],
            [("plan_id", "plan_id")],
        )
        row["result_status"] = row_formula(
            list(row),
            row_index,
            "result_status",
            'IF(IF({direction}="HIGHER_IS_BETTER",{actual_value}>={target_value},{actual_value}<={target_value}),"ON_TARGET","BELOW_TARGET")',
            row["result_status"],
        )
    actions = action_sheet(
        [
            action(
                "ACT-KPI-01",
                result.name,
                results[2]["kpi_result_id"],
                "UNIT-CS",
                "TEAM-CS",
                "RESPONSE_HOURS",
                value(results[2]["actual_value"]),
                24,
                "HOUR",
                "Rata-rata respons 30 jam dibanding target contoh maksimal 24 jam.",
                "Antrean layanan masih hipotesis; contoh hanya 3 pengukuran.",
                "Verifikasi sampel dan proses antrean, uji alokasi petugas, lalu ukur respons dengan metode dan cakupan yang sama.",
            )
        ]
    )
    return [indicators, plan, observation, result, actions], ["Master_Unit", "Master_Tim"]


SCENARIOS = [
    dict(
        id="keluhan",
        filename="01_Keluhan_Pelanggan.xlsx",
        title="Data keluhan pelanggan",
        builder=complaints,
        purpose="Keluhan anonim, SLA respons, status penyelesaian dan rekap produk/periode.",
        caution="Keluhan berulang bukan akar penyebab pasti. Pisahkan keluhan OPEN dari RESOLVED; jangan menghitung detail dan rekap bersama.",
        questions=[
            "Berapa keluhan dan pelanggaran SLA per produk bulan Januari 2026?",
            "Bagaimana perubahan waktu respons antara Januari dan Februari?",
        ],
    ),
    dict(
        id="marketing",
        filename="02_Marketing_ATL_BTL.xlsx",
        title="Data sebaran marketing ATL/BTL",
        builder=marketing,
        purpose="Campaign, channel, area, budget, spend, leads, biaya dan atribusi.",
        caution="Reach lintas channel dapat overlap. ATL tidak teratribusi memakai NULL, bukan nol; CPL/ROAS gabungan bukan rata-rata rasio dan bukan bukti efek kausal.",
        questions=[
            "Berapa biaya marketing menurut channel dan area periode Januari 2026?",
            "Bandingkan qualified leads BTL pada attribution model yang sama.",
        ],
    ),
    dict(
        id="kepuasan",
        filename="03_Kepuasan_Pelanggan_Produk.xlsx",
        title="Data kepuasan pelanggan berbasis produk",
        builder=satisfaction,
        purpose="Survei anonim 1-5, respons valid/tidak lengkap, penyebut CSAT dan rekap produk.",
        caution="CSAT memakai skor 4/5 dibagi respons valid, bukan NPS. Sampel kecil opt-in tidak memberi klaim representatif atau confidence statistik otomatis.",
        questions=[
            "Berapa respons valid dan respons puas per produk pada Januari 2026?",
            "Apakah perubahan CSAT memakai versi pertanyaan dan cakupan yang sebanding?",
        ],
    ),
    dict(
        id="competitiveness",
        filename="04_Competitiveness_Harga_Layanan.xlsx",
        title="Data competitiveness harga dan layanan",
        builder=competitiveness,
        purpose="Harga kemasan dinormalisasi per liter, waktu observasi dan skor produk/layanan/nilai.",
        caution="Produk, ukuran, promo, outlet, area dan waktu harus sebanding. Harga titik observasi berbeda dari skor survei periode; skor/bobot contoh bukan ranking pasar nyata.",
        questions=[
            "Bandingkan harga per liter pada produk/outlet/area/waktu setara.",
            "Bandingkan skor layanan per brand dengan versi survei dan bobot yang sama.",
        ],
    ),
    dict(
        id="kpi",
        filename="05_Plan_Goal_Pencapaian_KPI.xlsx",
        title="Data plan, goal dan pencapaian KPI",
        builder=kpi,
        purpose="Goal, indikator, target, bukti pengukuran, arah KPI, gap dan evaluasi tindakan.",
        caution="HIGHER dan LOWER memakai rumus berbeda; IDR, kasus, jam dan skor tidak boleh dijumlahkan bersama. Sampel 3 pengukuran hanya demonstrasi.",
        questions=[
            "Indikator mana yang belum mencapai target Januari 2026?",
            "Apa bukti pengukuran dan rencana tindak lanjut setiap plan?",
        ],
    ),
]


def sheet_collection(scenario):
    facts, master_names = scenario["builder"]()
    masters = refs()
    if not any(spec.name == "Master_Indikator" for spec in facts):
        facts = [indicator_master(), *facts]
    return [masters[name] for name in master_names] + facts


def deterministic_zip(entries):
    stream = io.BytesIO()
    # Stored entries avoid zlib-version differences between Windows and Linux builds.
    with ZipFile(stream, "w", ZIP_STORED) as archive:
        for name, content in sorted(entries.items()):
            info = ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_STORED
            archive.writestr(info, content)
    return stream.getvalue()


def workbook_bytes(scenario, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "CONTOH FIKTIF - " + scenario["title"]
    wb.properties.creator = "Google Sheet AI - contoh offline"
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    guide = wb.create_sheet("BACA_DULU")
    guide.append(["bagian", "petunjuk"])
    for row in [
        (
            "Peringatan",
            "SELURUH DATA DAN TARGET FIKTIF. Bukan data production, rekomendasi final atau bukti kausal.",
        ),
        ("Topik", scenario["purpose"]),
        ("Batas interpretasi", scenario["caution"]),
        (
            "Impor Google Sheets",
            "Upload Excel ke Drive lalu buka/konversi sebagai Google Sheets. Periksa formula, tanggal dan jumlah baris sebelum onboarding.",
        ),
        (
            "Waktu",
            "Tanggal disimpan sebagai teks ISO yyyy-mm-dd: pilih date, parse_date_id, format TEXT, pola %Y-%m-%d. Timestamp teks ISO ber-offset: timestamptz.",
        ),
        (
            "Master",
            "Master bersama identik pada beberapa file: muat sekali ke Registry master, lalu gunakan kembali binding. Jangan menggandakan registry master yang sama.",
        ),
        (
            "Rekap",
            "Rekap/rataan/rasio utama ditautkan ke raw lewat formula dan cache. Recalculate setelah edit/impor, periksa jumlah baris, penyebut dan master. CSV berisi hasil numerik, bukan formula.",
        ),
        (
            "Klasifikasi",
            "Master_* = MASTER. Tab data/rekap/Action_Plan = NON_MASTER. BACA_DULU, Kamus_Data, Rencana_ETL = petunjuk, Dikecualikan dari ETL.",
        ),
        (
            "Action plan",
            "Baseline diambil dari record sumber; cause_status UNVERIFIED, follow-up kosong/NOT_EVALUATED. Verifikasi bukti dan pilih target bisnis sebelum tindakan.",
        ),
        (
            "Load",
            "Contoh memakai key unik dan saran UPSERT. Review grain, reference binding, tipe, DQ, metrik, access dan rilis sebelum deploy. Jangan upload contoh ke sumber production yang aktif.",
        ),
    ]:
        guide.append(row)
    dictionary = wb.create_sheet("Kamus_Data")
    dictionary.append(["sheet", "column", "type", "nullable", "business_key", "meaning", "example"])
    plan = wb.create_sheet("Rencana_ETL")
    plan.append(["sheet", "classification", "grain", "business_key", "load", "metrics", "notes"])
    formula_cache = defaultdict(dict)
    for master in sheets:
        if master.role == "MASTER":
            wb.defined_names.add(
                DefinedName("REF_" + master.key, attr_text=f"'{master.name}'!$A$2:$A${len(master.rows) + 1}")
            )
    for spec in sheets:
        ws = wb.create_sheet(spec.name)
        ws.append(spec.headers)
        for row_number, row in enumerate(spec.rows, 2):
            ws.append([item.expression if isinstance(item, Formula) else item for item in row.values()])
            for index, (header, item) in enumerate(row.items(), 1):
                cell = ws.cell(row_number, index)
                if isinstance(item, Formula):
                    formula_cache[ws.title][cell.coordinate] = item.value
                if spec.kind(header) in ("text", "date", "timestamptz"):
                    cell.number_format = "@"
                elif "ratio" in header or header == "response_rate":
                    cell.number_format = "0.00%"
                elif spec.kind(header) == "numeric":
                    cell.number_format = "#,##0.00"
        table = Table(
            displayName="T_" + spec.name, ref=f"A1:{get_column_letter(len(spec.headers))}{ws.max_row}"
        )
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        for column_index, header in enumerate(spec.headers, 1):
            samples = [value(r[header]) for r in spec.rows if value(r[header]) is not None]
            nullable = any(value(r[header]) is None for r in spec.rows)
            explanation = (
                "Identitas stabil sesuai grain."
                if header == spec.key
                else "Waktu pengambilan, bukan periode bisnis."
                if header == "collected_at"
                else {
                    "period_start": "Tanggal pertama periode bisnis yang dicakup; teks ISO yyyy-mm-dd.",
                    "period_end": "Tanggal terakhir periode bisnis yang dicakup, inklusif; bukan tanggal pengambilan.",
                    "data_as_of_at": "Batas kondisi data yang berlaku; timestamp teks ISO dengan offset.",
                    "reported_at": "Waktu keluhan diterima; bukan waktu sync.",
                    "first_response_at": "Waktu respons awal diberikan kepada pelanggan.",
                    "resolved_at": "Waktu keluhan selesai; kosong jika masih OPEN.",
                    "first_response_hours": "Selisih respons awal - waktu laporan, dalam jam.",
                    "resolution_hours": "Selisih penyelesaian - laporan, dalam jam; kosong jika belum selesai.",
                    "response_breach_count": "1 jika respons awal melewati SLA contoh 24 jam, selainnya 0.",
                    "complaint_count": "1 untuk tiap keluhan unik; SUM menghasilkan jumlah keluhan.",
                    "resolved_count": "1 jika terselesaikan pada kondisi data yang dicatat, selainnya 0.",
                    "budget_idr": "Anggaran contoh dalam rupiah; bukan pengeluaran aktual.",
                    "spend_idr": "Pengeluaran aktual contoh dalam rupiah pada grain campaign/periode.",
                    "reach_count": "Reach dilaporkan/diestimasi pada channel; jangan SUM menjadi unique reach lintas channel.",
                    "leads_count": "Leads unik teratribusi menurut metode; NULL jika tidak terukur.",
                    "qualified_leads_count": "Bagian dari leads yang memenuhi definisi lead berkualitas.",
                    "attributed_revenue_idr": "Pendapatan yang teratribusi pada campaign dengan model yang dicatat; bukan uplift kausal.",
                    "cpl_idr": "spend_idr / leads_count; NULL jika penyebut nol/tidak diketahui.",
                    "roas_ratio": "attributed_revenue_idr / spend_idr; dapat lebih dari 1, bukan persentase konversi.",
                    "csat_ratio": "Respons dengan skor produk 4 atau 5 / respons valid; skala 0 sampai 1.",
                    "response_rate": "received_count / invited_count; bukan tingkat kepuasan.",
                    "invited_count": "Jumlah undangan survei dari log undangan yang sah; perbarui bila cakupan berubah.",
                    "is_valid_response": "TRUE bila respons memenuhi syarat metode; respons tidak lengkap dikecualikan dari CSAT.",
                    "sample_count": "Jumlah bukti/respons yang menjadi dasar angka; sampel kecil bukan bukti populasi.",
                    "weighted_score": "Skor produk*0.5 + layanan*0.3 + nilai*0.2, skala 1 sampai 5.",
                    "package_volume_ml": "Isi kemasan dalam ml, harus lebih dari nol.",
                    "normalized_price_idr_per_l": "Harga kemasan * 1000 / volume_ml untuk perbandingan per liter.",
                    "observed_at": "Waktu harga benar-benar diamati; bukan harga rata-rata periode.",
                    "target_value": "Target yang disetujui pada satuan dan arah indikator; target contoh harus diganti/direview.",
                    "actual_value": "SUM/AVG bukti pengukuran mengikuti definisi indikator; jangan gabungkan satuan berbeda.",
                    "direction": "HIGHER_IS_BETTER atau LOWER_IS_BETTER menentukan arah evaluasi.",
                    "unfavourable_gap": "Kekurangan/kelebihan yang merugikan menurut arah KPI; minimal 0.",
                    "achievement_ratio": "HIGHER: actual/target. LOWER: target/actual. NULL jika penyebut tidak positif; dapat >1.",
                    "baseline_value": "Nilai awal yang dibekukan pada periode bukti sebelum tindakan.",
                    "cause_status": "UNVERIFIED menunjukkan hipotesis belum menjadi akar penyebab yang terbukti.",
                    "evidence_ref": "Referensi bukti simulasi; ganti dengan bukti nyata yang aksesnya sesuai.",
                    "followup_value": "Hasil ukur setelah tindakan; NULL jika belum dievaluasi, bukan 0.",
                    "due_date": "Tenggat tindakan bisnis, bukan periode pengukuran.",
                    "verification_status": "NOT_EVALUATED berarti tindakan/hasil belum diverifikasi.",
                }.get(
                    header,
                    "Kode/label/atribut sesuai nama kolom dan grain; lihat Rencana_ETL serta definisi metrik. Jangan menjumlahkan rasio atau satuan berbeda.",
                )
            )
            dictionary.append(
                [
                    spec.name,
                    header,
                    spec.kind(header),
                    nullable,
                    header == spec.key,
                    explanation,
                    str(samples[0]) if samples else "Kosong: belum diukur/tidak berlaku; bukan 0",
                ]
            )
            if header.endswith("score") and spec.kind(header) in ("bigint", "numeric"):
                validation = DataValidation(
                    type="decimal", operator="between", formula1=1, formula2=5, allow_blank=nullable
                )
                validation.errorTitle = "Skor di luar skala"
                validation.error = "Gunakan skor 1 sampai 5 atau kosong jika tidak valid."
                validation.showErrorMessage = True
                ws.add_data_validation(validation)
                validation.add(f"{get_column_letter(column_index)}2:{get_column_letter(column_index)}1001")
            if "REF_" + header in wb.defined_names:
                validation = DataValidation(type="list", formula1="=REF_" + header, allow_blank=nullable)
                validation.errorTitle = "Kode belum ada di master"
                validation.error = "Pilih kode rujukan yang sah; master baru harus direview sebelum dipakai."
                validation.showErrorMessage = True
                ws.add_data_validation(validation)
                validation.add(f"{get_column_letter(column_index)}2:{get_column_letter(column_index)}1001")
        plan.append(
            [
                spec.name,
                spec.role,
                spec.grain,
                spec.key,
                "Registry master" if spec.role == "MASTER" else "UPSERT",
                "; ".join(
                    [
                        *(f"{m[0]}:{m[1]} [{m[3]}]" for m in spec.metrics),
                        *(
                            f"{m['code']}:{m['aggregation']} [{m['unit']}] dengan filter indikator"
                            for m in spec.scoped_metrics
                        ),
                    ]
                )
                or "Review definisi sebelum menambahkan metrik",
                spec.notes,
            ]
        )
        ws.sheet_properties.tabColor = "047857" if spec.role == "MASTER" else "2563EB"
    for ws in wb:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="065F46")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 32
        for index, header in enumerate([c.value for c in ws[1]], 1):
            width = (
                55 if header in ("petunjuk", "notes", "meaning") else min(38, max(18, len(str(header)) + 3))
            )
            ws.column_dimensions[get_column_letter(index)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide.column_dimensions["B"].width = 105
    wb.active = 0
    buffer = io.BytesIO()
    wb.save(buffer)
    with ZipFile(io.BytesIO(buffer.getvalue())) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    for sheet_index, ws in enumerate(wb, 1):
        name = f"xl/worksheets/sheet{sheet_index}.xml"
        if ws.title not in formula_cache:
            continue
        tree = ET.fromstring(entries[name])
        for cell in tree.findall(".//s:c", NS):
            if cell.attrib.get("r") not in formula_cache[ws.title]:
                continue
            cached = formula_cache[ws.title][cell.attrib["r"]]
            element = cell.find("s:v", NS)
            if element is None:
                element = ET.SubElement(cell, "{" + NS["s"] + "}v")
            if cached is None:
                cell.set("t", "str")
                element.text = None
            elif isinstance(cached, str):
                cell.set("t", "str")
                element.text = cached
            else:
                cell.attrib.pop("t", None)
                element.text = str(cached)
        entries[name] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    core = ET.fromstring(entries["docProps/core.xml"])
    for name in ("created", "modified"):
        node = core.find("{http://purl.org/dc/terms/}" + name)
        if node is not None:
            node.text = "2026-03-03T02:00:00Z"
    entries["docProps/core.xml"] = ET.tostring(core, encoding="utf-8", xml_declaration=True)
    return deterministic_zip(entries)


def csv_bytes(spec):
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(spec.headers)
    writer.writerows(
        [value(item) if value(item) is not None else "" for item in row.values()] for row in spec.rows
    )
    return buffer.getvalue().encode("utf-8-sig")


def build_documents():
    result, catalog = {}, []
    for index, scenario in enumerate(SCENARIOS, 29):
        sheets = sheet_collection(scenario)
        data = workbook_bytes(scenario, sheets)
        result[OUTPUT / scenario["filename"]] = data
        for spec in sheets:
            result[OUTPUT / "csv" / scenario["id"] / f"{spec.name}.csv"] = csv_bytes(spec)
            if spec.role == "NON_MASTER":
                result[OUTPUT / "mapping" / scenario["id"] / f"{spec.name}.json"] = (
                    json.dumps(spec.config(scenario["id"]), ensure_ascii=False, indent=2) + "\n"
                ).encode()
        facts = [spec for spec in sheets if spec.role == "NON_MASTER"]
        entry = dict(
            id=scenario["id"],
            title=scenario["title"],
            filename=scenario["filename"],
            download_url=f"/examples/analytics/{scenario['filename']}",
            description=scenario["purpose"],
            caution=scenario["caution"],
            questions=scenario["questions"],
            sheets=[
                dict(
                    name=spec.name,
                    rows=len(spec.rows),
                    role=spec.role,
                    grain=spec.grain,
                    business_key=spec.key,
                )
                for spec in sheets
            ],
            sha256=hashlib.sha256(data).hexdigest(),
        )
        catalog.append(entry)
        lines = [
            "SELURUH DATA FIKTIF untuk latihan; target/bobot bukan rekomendasi bisnis final.",
            scenario["purpose"],
            "",
            "| Tab | Grain | Key | Record contoh |",
            "| --- | --- | --- | --- |",
        ]
        lines.extend(f"| {spec.name} | {spec.grain} | {spec.key} | {len(spec.rows)} |" for spec in facts)
        lines += [
            "",
            "Unduh melalui Bantuan > Contoh Excel untuk analitik dan action plan atau /guide#contoh-excel.",
            f"File: {scenario['filename']}. BACA_DULU, Kamus_Data dan Rencana_ETL adalah petunjuk: Dikecualikan dari ETL.",
            "Master bersama dimuat sekali ke Registry master, lalu binding dipakai ulang. Data contoh memiliki kode anonim, tanggal teks ISO dan timestamp ber-offset; review sensitivitas data nyata.",
            "Tanggal: date + parse_date_id + TEXT + %Y-%m-%d. Timestamp: timestamptz. Key contoh unik, saran UPSERT. Metrik/rule/referensi/akses wajib direview; contoh mapping JSON bukan konfigurasi yang otomatis disetujui.",
            "Rasio/rekap utama memakai formula tertaut dan cache. Setelah raw/target diubah atau ditambah, recalculasi di Excel/Sheets dan validasi penyebut, cakupan dan key sebelum memakainya. CSV berisi nilai hasil, bukan formula.",
            "Waktu kejadian/pengukuran dan period_start/end berbeda dari collected_at/data_as_of_at. Jangan mengganti periode bisnis dengan tanggal sync.",
            "Action_Plan merujuk source_sheet/source_record_id, periode, baseline, owner dan tenggat. Cause_status UNVERIFIED: hipotesis perlu bukti. Followup_value kosong/NOT_EVALUATED bukan hasil 0.",
            "Action plan bukan engine penetapan tindakan otomatis. Verifikasi sampel/metode/causal evidence dan target dengan pemilik bisnis sebelum dijalankan.",
            "Batas interpretasi: " + scenario["caution"],
            "Pertanyaan contoh: " + " ".join(scenario["questions"]),
        ]
        body = "\n\n".join(["\n".join(lines[:2]), "\n".join(lines[3:])]) + "\n"
        if len(body) > 6000:
            raise ValueError("Knowledge article too long")
        header = (
            f"---\nid: contoh-excel-{scenario['id']}\ntitle: {json.dumps('Contoh Excel - ' + scenario['title'])}\n"
            f"summary: {json.dumps(scenario['purpose'])}\n"
            'routes: ["/guide", "/workspace"]\naudiences: ["*"]\nsource: docs/CONTOH_DATA_ANALITIK.md\n---\n'
        )
        result[ROOT / f"docs/knowledge/{index:02}-excel-{scenario['id']}.md"] = (header + body).encode()
    doc = [
        "# Lima contoh Excel untuk analitik dan action plan",
        "",
        "Semua data, target, harga, bobot dan hasil adalah fiktif. Gunakan sebagai acuan struktur dan evaluasi, bukan keputusan bisnis final.",
        "",
        "Unduh paket melalui **Bantuan > Contoh Excel untuk analitik dan action plan** atau **Panduan lengkap > Contoh Excel**.",
        "",
        "| Contoh | File Excel | Isi | Batas interpretasi |",
        "| --- | --- | --- | --- |",
    ]
    doc.extend(
        f"| {e['title']} | [Unduh](https://google.kanjabung.web.id/examples/analytics/{e['filename']}) | {e['description']} | {e['caution']} |"
        for e in catalog
    )
    doc += [
        "",
        "## Cara memakai",
        "",
        "1. Unduh Excel, baca BACA_DULU, Kamus_Data dan Rencana_ETL. Buat salinan untuk data nyata; jangan menimpa sumber production aktif dengan data contoh.",
        "2. Upload file ke Google Drive, buka/konversi sebagai Google Sheets. Periksa formula, format teks ISO, jumlah record dan nilai sebelum onboarding.",
        "3. Gunakan header baris 1 dan data baris 2. Tab petunjuk Dikecualikan. Master_* masuk alur Registry master; master identik antar-file dimuat sekali saja.",
        "4. Temukan tab, Profiling ulang dan klasifikasi. Mapping JSON di ZIP merupakan saran draft, bukan artifact approved. Terapkan tipe, DQ, master/taxonomy, key, metrik dan akses sesuai proses review aplikasi.",
        "5. Tanggal contoh adalah teks ISO: parse_date_id, TEXT, %Y-%m-%d. Timestamp sudah ber-offset +07:00: timestamptz. Rasio memakai 0..1, skor 1..5 dan uang IDR, bukan teks berisi unit.",
        "6. Rekap/rasio utama memakai formula tertaut ke raw/target dengan cache. Recalculasi setelah impor/edit, lalu periksa periode, key, penyebut dan jumlah record sebelum load. CSV berisi nilai contoh, tanpa formula; jangan memakai CSV lama setelah raw berubah.",
        "7. Review dan approval berbeda pembuat; deploy, sync, validasi record dan izin katalog memakai akun pembaca sebenarnya. Join hanya melalui relasi yang disetujui.",
        "8. Action_Plan membekukan baseline sesuai record/periode bukti. Verifikasi hipotesis, metrik/target, owner, tenggat, sampel dan hasil follow-up; jangan mengklaim tindakan sudah berhasil dari data yang masih kosong.",
        "",
        "## Prinsip evaluasi",
        "",
        "Jangan menjumlahkan detail dan rekap bersama, reach antar-channel, persen, nilai KPI berbeda satuan atau beberapa revisi dari record yang sama. Jangan AVG rasio tanpa bobot penyebut yang benar. Skor kepuasan tidak otomatis menjadi NPS; harga kemasan berbeda harus dinormalisasi. Sampel kecil/opt-in tidak membuktikan representasi populasi atau efek kausal tindakan.",
        "",
        "Data periodik memuat batas awal/akhir; event/survei memuat waktu bisnis; collected_at/data_as_of_at memberikan konteks pengambilan/kondisi data. Sel kosong berarti belum terukur/tidak berlaku, bukan nol.",
        "",
        "## Pemeliharaan",
        "",
        "Jalankan `python scripts/build_analytics_examples.py --frontend-root D:/projek/vue-googlesheet-ai`; opsi `--check` memastikan workbook, CSV, mapping, knowledge base dan helper tetap sinkron.",
    ]
    result[ROOT / "docs/CONTOH_DATA_ANALITIK.md"] = ("\n".join(doc) + "\n").encode()
    result[OUTPUT / "BACA_DULU.md"] = result[ROOT / "docs/CONTOH_DATA_ANALITIK.md"]
    result[ROOT / "docs/analytics-examples.json"] = (
        json.dumps(
            dict(examples=catalog, bundle_url="/examples/analytics/Paket_5_Contoh_Analitik.zip"),
            ensure_ascii=False,
            indent=2,
        )
        + "\n"
    ).encode()
    bundle_entries = {
        str(path.relative_to(OUTPUT)).replace("\\", "/"): content
        for path, content in result.items()
        if path.is_relative_to(OUTPUT)
    }
    result[OUTPUT / "Paket_5_Contoh_Analitik.zip"] = deterministic_zip(bundle_entries)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--frontend-root", type=Path)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    outputs = build_documents()
    if args.frontend_root:
        outputs[args.frontend_root / "docs/CONTOH_DATA_ANALITIK.md"] = outputs[
            ROOT / "docs/CONTOH_DATA_ANALITIK.md"
        ]
        outputs[args.frontend_root / "src/content/analyticsExamples.json"] = outputs[
            ROOT / "docs/analytics-examples.json"
        ]
        for path, content in list(outputs.items()):
            if path.parent == OUTPUT and path.suffix in (".xlsx", ".zip"):
                outputs[args.frontend_root / "public/examples/analytics" / path.name] = content
    for path, content in outputs.items():
        if args.check:
            if not path.exists() or path.read_bytes() != content:
                raise SystemExit(f"Example is stale: {path}")
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
    print(f"Verified {len(outputs)} example artifacts for five scenarios.")


if __name__ == "__main__":
    main()
