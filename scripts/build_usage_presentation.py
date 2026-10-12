"""Build an editable Indonesian training deck from the application's documented workflow."""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations"
OUT.mkdir(parents=True, exist_ok=True)
FILE = OUT / "Panduan_Penggunaan_Google_Sheet_AI.pptx"
W, H = 13.333333, 7.5
GREEN, INK, MUTED = "075E4B", "173B36", "617B75"
BG, LIGHT, WHITE, GOLD, BLUE = "F7FAF8", "E7F2ED", "FFFFFF", "B77912", "286B9D"
FONT = "Calibri"
APP = "https://google.kanjabung.web.id"
TEMPLATES = APP + "/examples/analytics/Paket_5_Contoh_Analitik.zip"
RULES_GUIDE = APP + "/examples/analytics/PANDUAN_ATURAN_ANALITIK.md"
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = "Dari Google Sheets ke Analitik dan Action Plan"
prs.core_properties.subject = "Panduan step by step Google Sheet AI untuk pengguna bisnis"
prs.core_properties.author = "Google Sheet AI • Panduan aplikasi"
prs.core_properties.keywords = "ETL, Google Sheets, kualitas data, analitik, action plan"
manifest = []


def rgb(color):
    return RGBColor.from_string(color)


def shape(slide, x, y, w, h, fill, kind=MSO_SHAPE.RECTANGLE, line=None):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = rgb(fill)
    item.line.fill.background()
    if line:
        item.line.color.rgb = rgb(line)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        item.adjustments[0] = 0.12
    return item


def text(slide, x, y, w, h, content, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT,
         font=FONT, margin=0, link=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(content.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(7)
        paragraph.line_spacing = 1.08
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        if link:
            run.hyperlink.address = link
    return box


def base(title, subtitle="", section="PANDUAN PRAKTIS", source="USER_GUIDE_END_TO_END.md", notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    shape(slide, 0, 0, W, 0.12, GREEN)
    text(slide, 0.65, 0.40, 11.6, 0.28, section, 12, GREEN, True)
    text(slide, 0.65, 0.94, 12.0, 0.63, title, 31, INK, True)
    if subtitle:
        text(slide, 0.65, 1.64, 12.0, 0.52, subtitle, 17, MUTED)
    shape(slide, 0.65, 6.97, 12.0, 0.012, "D8E5DF")
    text(slide, 0.65, 7.09, 10.9, 0.23, "Google Sheet AI  •  12 Oktober 2026  •  " + source, 10, MUTED)
    text(slide, 11.8, 7.04, 0.85, 0.32, f"{len(prs.slides):02d}", 14, GREEN, True, PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = (
        notes + "\n\nRujukan aplikasi: docs/" + source +
        "\nMateri sesuai dokumentasi dan implementasi per 12 Oktober 2026. "
        "Semua angka ilustrasi dalam presentasi adalah simulasi, bukan data production."
    )
    manifest.append({"slide": len(prs.slides), "title": title, "source": source, "speaker_notes": notes})
    return slide


def card(slide, x, y, w, h, title, body, accent=GREEN, size=18, number=None):
    shape(slide, x, y, w, h, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
    shape(slide, x, y, 0.05, h, accent)
    tx = x + 0.23
    if number is not None:
        shape(slide, tx, y + 0.20, 0.45, 0.45, accent, MSO_SHAPE.OVAL)
        text(slide, tx, y + 0.25, 0.45, 0.30, str(number), 16, WHITE, True, PP_ALIGN.CENTER)
        tx += 0.62
    text(slide, tx, y + 0.20, w - (tx - x) - 0.22, 0.55, title, 21, accent, True)
    text(slide, x + 0.23, y + 0.89, w - 0.46, h - 1.02, body, size)


def two(slide, first, second, y=2.3, h=3.65):
    card(slide, 0.65, y, 5.85, h, *first)
    card(slide, 6.75, y, 5.85, h, *second)


def ribbon(slide, message, color=GREEN):
    shape(slide, 0.65, 6.25, 12, 0.46, LIGHT if color == GREEN else "FFF2D7", MSO_SHAPE.ROUNDED_RECTANGLE)
    text(slide, 0.84, 6.34, 11.6, 0.28, message, 15, color, True)


def table(slide, headers, rows, widths, x=0.65, y=2.32, h=3.63, font=17):
    item = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(sum(widths)), Inches(h))
    tab = item.table
    for column, width in zip(tab.columns, widths):
        column.width = Inches(width)
    for row_index, row in enumerate([headers, *rows]):
        for column_index, value in enumerate(row):
            cell = tab.cell(row_index, column_index)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(GREEN if row_index == 0 else (WHITE if row_index % 2 else LIGHT))
            cell.margin_left = cell.margin_right = Inches(0.14)
            cell.margin_top = cell.margin_bottom = Inches(0.1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                p.line_spacing = 1.05
                p.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = rgb(WHITE if row_index == 0 else INK)
    return tab


def link_button(slide, x, y, w, title, address):
    item = shape(slide, x, y, w, 0.61, GREEN, MSO_SHAPE.ROUNDED_RECTANGLE)
    item.click_action.hyperlink.address = address
    text(slide, x + 0.17, y + 0.15, w - 0.34, 0.35, title, 16, WHITE, True, link=address)


# 01 — Cover
s = base("", source="USER_GUIDE_END_TO_END.md", notes="Buka dengan kebutuhan bisnis: mengubah pencatatan yang tersebar menjadi informasi yang dapat dipercaya. Presentasi ini dapat digunakan untuk pelatihan 45–60 menit. Sesuaikan tempo dan sisipkan demo menggunakan data contoh, bukan sumber production aktif.")
s.background.fill.fore_color.rgb = rgb(GREEN)
text(s, 0.7, 0.72, 9, 0.45, "PANDUAN PENGGUNAAN • UNTUK TIM BISNIS", 15, "AEE0CA", True)
text(s, 0.7, 1.64, 8.2, 2.55, "Dari Google Sheets\nke Analitik\ndan Action Plan", 43, WHITE, True)
text(s, 0.73, 4.6, 7.9, 1.12, "Siapkan sumber yang rapi. Review definisinya.\nBaca hasilnya. Putuskan tindakan berdasarkan bukti.", 23, "D9EFE6")
for i, (heading, caption) in enumerate([("DATA", "Kolom • kode • periode"), ("ANALITIK", "Metrik • tabel • chart"), ("TINDAKAN", "Aturan • pemilik • evaluasi")]):
    shape(s, 9.15, 1.65 + i * 1.45, 3.4, 1.16, "0C745D", MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 9.4, 1.85 + i * 1.45, 2.9, 0.4, heading, 23, WHITE, True)
    text(s, 9.4, 2.38 + i * 1.45, 2.9, 0.3, caption, 15, "D9EFE6")
text(s, 0.73, 6.4, 9, 0.38, "Google Sheet AI  |  Panduan praktis  |  12 Oktober 2026", 15, "D9EFE6")

# 02 — Goals
s = base("Tujuan: keputusan yang bisa dijelaskan", "Ukuran keberhasilan dimulai dari mutu data dan definisi bisnis.", notes="Minta peserta menyebut satu keputusan yang ingin diperbaiki. Misalnya unit mana perlu perbaikan SLA. Jelaskan bahwa AI tidak menggantikan definisi indikator, data yang benar, atau review keputusan oleh pemilik bisnis.")
for x, y, title, body in [
    (0.65, 2.3, "Satu bahasa bisnis", "Nama produk, unit, kategori, satuan dan periode dipahami dengan arti yang sama."),
    (6.75, 2.3, "Angka bisa ditelusuri", "Jawaban memiliki sumber, rumus, cakupan periode dan status pemuatan."),
    (0.65, 4.3, "Akses sesuai kebutuhan", "Pembaca hanya melihat produk, kolom dan baris yang diizinkan."),
    (6.75, 4.3, "Perbaikan terukur", "Usulan tindakan punya pemilik, tenggat, baseline dan bukti evaluasi.")]:
    card(s, x, y, 5.85, 1.8, title, body, size=17)

# 03 — Outputs / current boundaries
s = base("Output yang tersedia di aplikasi", "Pisahkan hasil analisis, usulan tindakan, dan keputusan pelaksanaannya.", source="PANDUAN_ATURAN_ANALITIK.md", notes="Tabel/chart tersedia di Dashboard melalui query builder atau pertanyaan bahasa alami. Backend aturan action plan telah tersedia, tetapi impor, preview, approval dan evaluate saat ini dilakukan melalui API. Jangan menyebut adanya tombol action plan otomatis di wizard atau chat. Output action plan adalah usulan, bukan pelaksanaan.")
table(s, ["Output", "Isi yang dihasilkan", "Cara mengakses"], [
    ["Data product", "Data hasil ETL, definisi kolom, metrik dan sumber", "Katalog data"],
    ["Analitik", "Jawaban angka per produk / unit / periode", "Dashboard • query builder • pertanyaan AI"],
    ["Chart", "Visualisasi hasil query sesuai metrik dan dimensi", "Pilihan chart pada hasil analitik"],
    ["Action plan", "Usulan + aturan pemicu + nilai bukti + pemilik", "Backend API • approval terpisah"]
], [2.05, 5.85, 4.1], font=17)
ribbon(s, "Aturan action plan belum dijalankan otomatis oleh wizard ETL atau chat AI.", GOLD)

# 04 — Cycle
s = base("Alur ideal: data → analitik → perbaikan", "Pengukuran ulang menutup siklus; chart saja belum membuktikan perbaikan.", source="PANDUAN_FORMAT_GOOGLE_SHEET.md", notes="Tekankan bahwa action plan yang baik memerlukan baseline dan follow-up dengan definisi/metode sebanding. Korelasi atau tren menurun bukan otomatis bukti bahwa tindakan tertentu menyebabkan perubahan.")
labels = [("01", "Catat data", "Kejadian dan periode jelas"), ("02", "Validasi", "Key, tipe, referensi, akses"),
          ("03", "Analisis", "Angka dan chart\nPerbandingan"), ("04", "Tindak lanjut", "Pemilik dan tenggat"), ("05", "Ukur ulang", "Baseline vs hasil")]
for i, (num, title, body) in enumerate(labels):
    x = 0.65 + i * 2.43
    shape(s, x, 2.62, 2.23, 2.55, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.19, 2.9, 1.85, 0.6, num, 36, GREEN, True)
    text(s, x + 0.19, 3.63, 1.85, 0.48, title, 20, INK, True)
    text(s, x + 0.19, 4.31, 1.85, 0.69, body, 17, MUTED)
ribbon(s, "Hasil ideal: keputusan dapat ditelusuri dan dampak tindakan diperiksa dengan bukti.")

# 05 — Roles
s = base("Siapa melakukan apa?", "Persetujuan dilakukan pihak yang berwenang dan berbeda dari pembuat.", notes="Role teknis tidak otomatis memberikan hak membaca semua data. Admin menyiapkan assignment dan policy; approver harus ditunjuk untuk sumber. Pemilik bisnis memverifikasi target dan manfaat tindakan. Gate IT serta unit berlaku jika aturan siap tayang diaktifkan.")
table(s, ["Peran", "Tanggung jawab", "Hasil yang diperiksa"], [
    ["Admin platform", "Akun, assignment, policy dan approver", "Izin organisasi dan akses sumber"],
    ["Source Owner / Steward", "Siapkan Sheet, konfigurasi ETL dan sync", "Data, mapping dan kualitas"],
    ["Technical Approver", "Review konfigurasi / aturan; approve atau reject", "Versi, bukti, keamanan dan blocker"],
    ["Analyst / Viewer", "Membaca katalog, query dan chart sesuai izin", "Periode, metrik dan cakupan hasil"],
    ["Pemilik bisnis", "Menetapkan target; menilai usulan dan hasil tindakan", "Manfaat, penanggung jawab dan follow-up"]
], [2.7, 5.25, 4.05], font=16)

# 06 — End-to-end map
s = base("Peta penggunaan: dari sumber sampai output", "Gunakan urutan ini sebagai checklist saat onboarding dataset baru.", notes="Master dan taxonomy diperlukan bila ada rujukan atau kategori yang perlu dibakukan. Deploy menyiapkan struktur dan mengaktifkan konfigurasi; pemuatan data tetap tahap terpisah. Aturan action plan mempunyai alur approval tersendiri melalui API.")
steps = [("Siapkan Sheet", "Format + definisi bisnis"), ("Hubungkan", "URL + akses Viewer"), ("Temukan & profil", "Tab + kondisi data"),
         ("Klasifikasikan", "Master / data / kontrol"), ("Setup config", "7 langkah wizard"), ("Review", "Reviewer + gate akses"),
         ("Deploy", "Konfigurasi ACTIVE"), ("Sync & validasi", "Baris benar-benar dimuat"), ("Analisis", "Katalog + query + chart"), ("Action plan", "Aturan + evaluasi API")]
for i, (title, caption) in enumerate(steps):
    x, y = 0.65 + (i % 5) * 2.43, 2.4 + (i // 5) * 1.77
    shape(s, x, y, 2.22, 1.57, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.18, y + 0.12, 1.8, 0.37, f"{i+1:02d}", 22, GREEN, True)
    text(s, x + 0.18, y + 0.66, 1.87, 0.45, title, 18, INK, True)
    text(s, x + 0.18, y + 1.16, 1.87, 0.26, caption, 12, MUTED)

# 07 — Workbook tabs
s = base("Format file dan tab: nama harus bermakna", "Nama file + tab membantu saran nama bisnis dan produk; tetap periksa di konfigurasi.", section="BAGIAN 1 • SIAPKAN GOOGLE SHEETS", source="PANDUAN_FORMAT_GOOGLE_SHEET.md", notes="Contoh nama file: Keluhan Pelanggan KAN Jabung 2026. Contoh tab: Keluhan_Data. Nama yang dihasilkan tidak menggantikan deskripsi bisnis dan grain. Dua tab kontrol aturan/formula perlu dikecualikan dari ETL dan diimpor lewat API setelah produk utamanya aktif.")
table(s, ["Tab contoh", "Tujuan isi", "Klasifikasi / proses"], [
    ["Master_Produk / Master_Unit", "Kode rujukan baku, label dan definisi", "MASTER • Registry master"],
    ["Keluhan_Data / Survey_Data", "Catatan kejadian atau pengukuran", "NON_MASTER • ETL"],
    ["Keluhan_Rekap / KPI_Result", "Ringkasan dengan grain dan rumus jelas", "NON_MASTER • jangan hitung ganda"],
    ["Action_Plan", "Catatan pelaksanaan dan follow-up manusia", "NON_MASTER bila dimuat"],
    ["Formula_Analitik / Aturan_Action_Plan", "Definisi perhitungan dan pemicu usulan", "Dikecualikan dari ETL • API"],
    ["BACA_DULU / Kamus_Data / Rencana_ETL", "Petunjuk dan arti setiap field", "Dikecualikan dari ETL"]
], [3.8, 4.75, 3.45], font=15)

# 08 — Grain and table
s = base("Satu baris = satu catatan yang jelas", "Contoh grain: satu baris per keluhan, bukan gabungan keluhan dan subtotal.", source="PANDUAN_FORMAT_GOOGLE_SHEET.md", notes="Business key complaint_id harus unik dan stabil. Contoh berikut adalah ilustrasi struktur, bukan cuplikan production. Gunakan header tunggal dan unik; hindari merged cells, judul menumpuk, baris kosong serta footer total di range data. Rekap dapat berada pada tab terpisah dengan grain berbeda.")
table(s, ["complaint_id", "event_date", "unit_code", "product_code", "status"], [
    ["KLH-001", "2026-01-05", "UNIT-CS", "PROD-MILK", "OPEN"],
    ["KLH-002", "2026-01-06", "UNIT-CS", "PROD-FEED", "RESOLVED"],
    ["KLH-003", "2026-01-07", "UNIT-CS", "PROD-MILK", "OPEN"]
], [2.35, 2.35, 2.35, 2.65, 2.3], h=2.8, font=18)
text(s, 0.8, 5.48, 11.55, 0.65, "Header tunggal • tanpa merge • key unik • tidak ada subtotal di tengah data", 22, GREEN, True)
ribbon(s, "Ilustrasi struktur. Key kosong / duplikat harus diperbaiki sebelum data digunakan.")

# 09 — Field rationale
s = base("Mengapa kolom-kolom ini perlu ada?", "Isi Sheet menentukan apakah analisis bisa menjawab pertanyaan bisnis dengan benar.", source="PANDUAN_FORMAT_GOOGLE_SHEET.md", notes="Kebutuhan kolom bergantung pada grain. Jangan menambah tanggal kejadian fiktif ke data rekap. Untuk data periodik, awal/akhir periode bisnis wajib eksplisit. Identitas atau kontak pribadi harus dilindungi dan tidak dijadikan dimensi analitik sensitif.")
table(s, ["Informasi", "Contoh field", "Tujuan"], [
    ["Identitas catatan", "complaint_id / plan_id", "Mencegah hitung ganda dan menautkan bukti"],
    ["Konteks bisnis", "unit_code / product_code / category_code", "Membandingkan kelompok dengan arti konsisten"],
    ["Waktu bisnis", "event_at / period_start / period_end", "Menjawab kapan kejadian atau kinerja terjadi"],
    ["Ukuran dan satuan", "actual_value / measurement_unit", "Menghitung ukuran yang sebanding"],
    ["Target dan arah", "target_value / direction", "Menilai capaian sesuai tujuan indikator"],
    ["Basis pengukuran", "sample_count / denominator / method", "Menilai cakupan dan keterbatasan angka"],
    ["Jejak pengambilan", "collected_at / data_as_of_at", "Mengetahui umur dan asal bukti"]
], [2.3, 4.45, 5.25], font=15)

# 10 — Date meaning
s = base("Periode bisnis berbeda dari waktu sync", "Data Januari tetap kinerja Januari meskipun dikumpulkan pada Februari.", source="PANDUAN_FORMAT_GOOGLE_SHEET.md", notes="period_start 2026-01-01 dan period_end 2026-01-31 mencakup Januari. collected_at 2026-02-02T09:00:00+07:00 adalah waktu pengambilan. Jika source berupa teks ISO, pilih date + parse_date_id + TEXT + %Y-%m-%d. Jika native tanggal Sheets dibaca sebagai serial, pilih format Serial Google Sheets; 46023 = 2026-01-01. Bigint biasa tidak boleh diasumsikan serial tanggal.")
two(s, ("Apa yang diukur?", "Periode kinerja\n2026-01-01 sampai 2026-01-31\n\nKejadian / pengukuran\nevent_at atau measured_at"),
    ("Kapan data dikumpulkan?", "collected_at\n2026-02-02T09:00:00+07:00\n\nWaktu sync\nWaktu pemuatan oleh aplikasi"), h=3.65)
ribbon(s, "Tanggal teks ISO → target date + parse_date_id + TEXT + pola %Y-%m-%d.")

# 11 — Controlled vocabulary
s = base("Kode baku membuat istilah konsisten", "Master adalah daftar rujukan; taxonomy adalah daftar kategori yang disepakati.", source="ISTILAH_TEKNIS.md", notes="Jelaskan perbedaan master produk/unit dengan taxonomy jenis keluhan atau prioritas. Dropdown di Sheets membantu input, tetapi tidak menggantikan binding dan validasi aplikasi. Alias harus diverifikasi; AI tidak otomatis berhak memilih relasi yang ambigu.")
two(s, ("Master data", "PROD-MILK → Produk susu\nUNIT-CS → Unit layanan\nTEAM-CS → Tim layanan\n\nKode stabil; label boleh diperjelas."),
    ("Taxonomy", "Kategori: PRODUCT / SERVICE\nStatus: OPEN / RESOLVED\nPrioritas: LOW / MEDIUM / HIGH\n\nGunakan daftar dan alias resmi."))
ribbon(s, "Gunakan dropdown; review definisi, record master, versi kategori dan binding sebelum analisis.")

# 12 — Precision and ratios
s = base("Rumus yang benar dimulai dari basis angka", "Jangan merata-ratakan persen jika ukuran kelompoknya berbeda.", source="PANDUAN_FORMAT_GOOGLE_SHEET.md", notes="Ilustrasi survei: A memiliki 9 puas dari 10 valid, B memiliki 9 puas dari 90 valid. Rata-rata 90% dan 10% adalah 50%, tetapi CSAT gabungan adalah 18/100 = 18%. Rumus persentase perlu pembilang dan penyebut, bukan hanya kolom persentase. Skor 1–5 dan rasio 0–1 merupakan skala yang berbeda.")
table(s, ["Kelompok", "Respons puas", "Respons valid", "CSAT"], [
    ["A", 9, 10, "90%"], ["B", 9, 90, "10%"], ["Gabungan", 18, 100, "18% = 18 / 100"]
], [2.9, 2.7, 2.8, 3.6], h=2.6, font=21)
text(s, 0.85, 5.28, 11.6, 0.76, "Simpan pembilang + penyebut • satuan jelas • NULL berbeda dari 0", 23, GREEN, True)
ribbon(s, "SIMULASI • Akurasi tidak dijamin oleh banyak digit; metode dan cakupan tetap diperiksa.", GOLD)

# 13 — Upload
s = base("Langkah 1 — Siapkan dan upload sumber", "Aplikasi menerima sumber Google Sheets melalui URL / ID spreadsheet.", section="BAGIAN 2 • HUBUNGKAN DAN KONFIGURASIKAN", notes="Jika sumber masih Excel, upload ke Google Drive lalu buka atau konversi menjadi Google Sheets. Periksa hasil formula, jumlah baris dan tanggal setelah konversi. Service account backend diberi Viewer menggunakan email yang disediakan admin. Aplikasi membaca dan tidak menulis ulang Sheet. Gunakan contoh fiktif saat latihan.")
two(s, ("Di Google Drive / Sheets", "1. Upload Excel jika diperlukan.\n2. Konversi ke Google Sheets.\n3. Cek header, formula, tanggal, key.\n4. Bagikan ke service account: Viewer."),
    ("Di aplikasi", "1. Login sesuai peran.\n2. Buka Sumber & tracking / Workspace.\n3. Gunakan URL / ID spreadsheet.\n4. Gunakan sumber lama jika sudah terdaftar."))
ribbon(s, "Lanjut jika: Sheet tersedia, bisa dibaca service account dan isi sumber sudah diperiksa.")

# 14 — Registration
s = base("Langkah 2 — Hubungkan & profiling", "Isi identitas sumber sekaligus tanggung jawab dan cakupan aksesnya.", notes="Metadata sumber mencakup unit pemilik, domain bisnis, yurisdiksi, purpose, data owner/steward dan sensitivitas. Opsi unit/domain/yurisdiksi mengikuti assignment aktif pengguna. Jika kosong, hubungi admin, bukan memilih scope sembarang. Hubungkan & profiling mendaftarkan sumber dan memulai job discovery/profiling. Pantau job sampai sukses; sumber dapat tetap tercatat meskipun job gagal.")
two(s, ("Isi formulir sumber", "Nama sumber + URL / ID\nUnit pemilik + domain bisnis\nYurisdiksi + tujuan penggunaan\nData owner + data steward\nSensitivitas data"),
    ("Jalankan dan periksa", "Klik Hubungkan & profiling.\nTunggu QUEUED / RUNNING selesai.\nBuka monitor jika FAILED.\nPastikan tab dan header terbaca.\nJangan daftarkan ulang sumber yang sama."))
ribbon(s, "Lanjut jika: discovery/profiling sukses; sumber dan daftar tab tersedia.")

# 15 — Discovery and classification
s = base("Langkah 3 — Pilih dan klasifikasikan tab", "Setiap tab mempunyai peran; tab petunjuk dan kontrol tidak dimuat sebagai data bisnis.", notes="Pilih MASTER untuk rujukan stabil dan NON_MASTER untuk kejadian, aktivitas, pengukuran atau rekap. Dikecualikan dipakai untuk petunjuk dan dua tab kontrol analitik. Setelah tab baru atau rename, jalankan Temukan tab / discovery ulang dan profiling ulang pada sumber yang sama. Profiling memberi gambaran awal, bukan jaminan seluruh record valid.")
table(s, ["Kegiatan", "Yang diperiksa", "Hasil"], [
    ["Temukan tab", "Tab baru, rename atau tab yang hilang", "Daftar tab diperbarui"],
    ["Profiling", "Header, dugaan tipe, kosong dan sensitivitas", "Bahan awal review mapping"],
    ["MASTER", "Unit, produk, indikator dan rujukan stabil", "Definisi / record / binding master"],
    ["NON_MASTER", "Transaksi, survei, pengukuran atau rekap", "Konfigurasi ETL per tab"],
    ["Dikecualikan", "Petunjuk, Formula_Analitik, Aturan_Action_Plan", "Tidak dimuat sebagai tabel ETL"]
], [2.55, 5.45, 4.0], font=16)
ribbon(s, "Lanjut jika: klasifikasi terkonfirmasi dan master / taxonomy yang diperlukan sudah siap.")

# 16 — Seven wizard steps
s = base("Langkah 4 — Setup config: tujuh bagian", "Konfigurasi menjelaskan arti data, cara mengolahnya dan siapa boleh membacanya.", source="PANDUAN_REVIEW_ETL.md", notes="Nama bagian mengikuti wizard aktual. Identitas dataset menentukan grain. Kolom dan sensitivitas menentukan mapping, tipe, nullable, key, PII dan referensi. Cleansing menentukan transformasi. Kualitas data menentukan aturan dan tindakan gagal. Pemuatan memilih target/strategi. Analitik dan akses menentukan dimensi/metrik/role. Validasi dan persetujuan menguji, meninjau perubahan schema dan menyelesaikan review.")
table(s, ["Bagian wizard", "Keputusan utama"], [
    ["1. Identitas dataset", "Nama bisnis, deskripsi dan arti satu baris"],
    ["2. Kolom dan sensitivitas", "Mapping, tipe, nullable, key, referensi, PII"],
    ["3. Cleansing", "Trim dasar dan transformasi tambahan yang diperlukan"],
    ["4. Kualitas data", "Aturan validasi dan tindakan ketika gagal"],
    ["5. Pemuatan", "Tabel tujuan dan APPEND / UPSERT / FULL_REFRESH"],
    ["6. Analitik dan akses", "Dimensi, metrik, unit, sinonim dan role pembaca"],
    ["7. Validasi dan persetujuan", "Dry-run, preview schema, checklist dan review"]
], [4.35, 7.65], font=17)

# 17 — Mapping
s = base("Mapping: dari header Sheet ke field tujuan", "Jenis isi sumber dan tipe tujuan harus cocok; boleh kosong adalah keputusan bisnis.", source="PANDUAN_REVIEW_ETL.md", notes="Contoh mapping adalah ilustrasi. complaint_id text wajib terisi dan key. event_date date adalah tanggal bisnis. Jika input merupakan ISO text gunakan parse_date_id/TEXT/%Y-%m-%d; jika serial gunakan pilihan Serial Google Sheets. completed_at boleh kosong ketika belum selesai; jangan membuat nol/tanggal palsu. Numeric merupakan pilihan angka uang/pengukuran. Kolom identitas pribadi perlu sensitivitas dan akses yang sesuai.")
table(s, ["Header sumber", "Field / tipe tujuan", "Aturan contoh"], [
    ["complaint_id", "complaint_id • text", "Wajib • key unik • kode tetap"],
    ["event_date", "event_date • date", "Wajib • format parsing sesuai sumber"],
    ["unit_code", "unit_code • text", "Wajib • rujuk master unit"],
    ["actual_value", "actual_value • numeric", "Satuan / presisi sesuai pengukuran"],
    ["completed_at", "completed_at • timestamptz", "Boleh kosong selama belum selesai"],
    ["Narasi keluhan", "Field text dengan sensitivitas sesuai isi", "Periksa PII sebelum dipublikasikan"]
], [3.1, 4.1, 4.8], font=16)
ribbon(s, "Header sumber harus cocok persis dengan profil terbaru; rename dilakukan di Sheet lalu profiling ulang.")

# 18 — Cleansing / DQ
s = base("Cleansing dan DQ menjaga arti data", "Rapikan nilai tanpa mengubah makna atau menutupi kesalahan sumber.", source="PANDUAN_REVIEW_ETL.md", notes="Trim otomatis hanya membuang spasi awal/akhir untuk text/varchar. Spasi di tengah, huruf dan nol depan dipertahankan. Spasi saja menjadi kosong/NULL dan tetap gagal bila kolom wajib. Tambahan transformasi harus direview. DQ dapat WARN, REJECT_ROW, STOP_BATCH atau REQUIRE_REVIEW. Jangan menganggap semua pemeriksaan lintas kolom atau overlap waktu otomatis tersedia.")
two(s, ("Cleansing", "\"  UNIT-CS  \" → \"UNIT-CS\"\nTrim otomatis pada text / varchar.\n\nTanggal: parsing sesuai format.\nKategori: referensi resmi.\nTransformasi tambahan direview."),
    ("Kualitas data", "Key tidak kosong / tidak berulang.\nTipe dan nilai wajib valid.\nKode ditemukan di master.\nSkor / nilai dalam batas yang sah.\nTemuan punya tindakan dan pemilik."))
ribbon(s, "Nullable mengizinkan kosong; bukan mengizinkan angka / tanggal yang formatnya salah.")

# 19 — Load strategy
s = base("Strategi pemuatan: pilih sesuai sumber", "Business key menentukan apakah perubahan memperbarui record atau menambah record.", source="PANDUAN_REVIEW_ETL.md", notes="APPEND untuk catatan baru; sync ulang kejadian lama tanpa kontrol bisa menggandakan data. UPSERT memperbarui berdasarkan business key stabil. FULL_REFRESH mengganti snapshot lengkap dan tetap menggunakan review batch, tidak jalur sync langsung NON_MASTER. Jangan menerapkan FULL_REFRESH pada sumber parsial.")
for i, (title, body) in enumerate([
    ("APPEND", "Tambahkan record baru.\n\nCocok: log kejadian baru.\nPeriksa risiko hitung ganda."),
    ("UPSERT", "Update jika key sudah ada; tambah jika belum.\n\nCocok: koreksi data dengan key stabil."),
    ("FULL_REFRESH", "Ganti isi dengan snapshot lengkap.\n\nCocok: snapshot menyeluruh.\nTetap melalui review batch.")]):
    card(s, 0.65 + i * 4.08, 2.4, 3.85, 3.65, title, body, size=20)
ribbon(s, "Untuk revisi dataset yang sama, pertahankan target tabel kecuali perpindahan sudah direncanakan.")

# 20 — Semantic
s = base("Analitik dan akses: beri AI konteks yang jelas", "AI memakai definisi produk dan metrik yang disetujui, bukan menebak arti kolom angka.", source="PANDUAN_FORMAT_GOOGLE_SHEET.md", notes="Draft manual dapat memiliki metrics kosong; tambahkan metrik sebelum berharap analitik angka bekerja. Contoh SUM complaint_count menghasilkan jumlah keluhan. Dimensi digunakan untuk kelompok/filter. Definisi, unit, sinonim dan periode membantu AI memilih metrik. Agregasi editor dibatasi sum/count/avg/min/max/count_distinct; formula gabungan menggunakan backend aturan atau hasil perhitungan Sheet yang direview. Kolom sensitif MEDIUM/HIGH tidak dijadikan dimensi/metrik.")
two(s, ("Definisi analitik", "Dimensi: produk, unit, periode.\nMetrik: jumlah keluhan = SUM.\nUnit: CASE, IDR, HOUR atau RATIO.\nSinonim: istilah yang digunakan tim.\nDeskripsi: basis dan batas perhitungan."),
    ("Akses dan konteks", "Role pembaca sesuai kebutuhan.\nPolicy sumber / produk berlaku.\nKolom dan baris mengikuti izin.\nPeriode dan metode pengukuran jelas.\nRelasi dataset memakai join approved."))
ribbon(s, "Nama bisnis + grain + definisi metrik + unit + periode = konteks analitik yang dapat diperiksa.")

# 21 — Approval / safe migration
s = base("Langkah 5 — Validasi, review dan approval", "Periksa bukti dan dampak sebelum konfigurasi ditayangkan.", source="PANDUAN_REVIEW_ETL.md", notes="Dry-run menampilkan baris valid/bermasalah, sebelum-sesudah dan blocker. Preview schema memuat kolom tambahan, tipe, nullability, key, dampak view/dependency dan konversi tanggal. Reviewer berbeda menilai mapping dan schema pada revisi yang sama. Approval gabungan dapat menyetujui metadata dan mengaktifkan policy SOURCE yang sudah APPROVED. Gate IT/unit tetap berlaku bila diwajibkan. Perubahan schema yang didukung dieksekusi worker saat Deploy dalam transaksi; perubahan yang tidak aman diblokir.")
two(s, ("Editor menyiapkan bukti", "Simpan draft dan jalankan dry-run.\nPeriksa mapping serta hasil konversi.\nTinjau preview perubahan schema.\nSelesaikan blocker dan checklist.\nAjukan review pada revisi yang sama."),
    ("Reviewer memberi keputusan", "Periksa grain, DQ dan sensitivitas.\nPeriksa akses / policy sumber.\nApprove atau reject dengan alasan.\nLengkapi gate IT / unit jika wajib.\nSumber berubah → review ulang."))
ribbon(s, "Worker memigrasikan perubahan schema yang didukung saat Deploy; blocker harus diselesaikan dahulu.")

# 22 — Deploy / sync distinct
s = base("Langkah 6 — Deploy, lalu muat data", "Konfigurasi ACTIVE belum membuktikan data sudah masuk ke tabel.", notes="Deploy menyiapkan/menyesuaikan struktur target dan semantic view. Jalankan Sync manual di Workspace atau Sumber & tracking. NON_MASTER ACTIVE yang seluruh validasi/gate-nya lolos dapat dimuat langsung; master, FULL_REFRESH, jadwal/batch biasa tetap preview/approval/Apply. Job induk sukses belum berarti semua tab sukses. Muat ulang hanya membaca status, tidak menjalankan sync.")
two(s, ("DEPLOY", "Jalankan Deploy konfigurasi.\nPantau job sampai SUCCEEDED.\n\nHasil: konfigurasi ACTIVE\ndan struktur / produk siap."),
    ("SYNC / PEMUATAN", "Klik Sync manual per sumber.\nPeriksa hasil setiap tab.\n\nHasil: baris dimuat, waktu pemuatan\ndan batch per tab SUCCEEDED."))
ribbon(s, "Master / FULL_REFRESH / batch biasa tetap membutuhkan preview, approval dan Apply.", GOLD)

# 23 — Errors
s = base("Jika proses berhenti, lihat alasan per tab", "Gunakan monitor job dan Lihat detail; perbaiki penyebab sebelum mencoba ulang.", source="ERROR_RESPONSES.md", notes="SOURCE_READ_FAILED memiliki beberapa kemungkinan termasuk nama tab/range; SOURCE_ACCESS_DENIED berkaitan izin service account. REVIEW_STALE berarti snapshot berubah sejak review. NEEDS_INPUT adalah status menunggu tindakan, bukan worker sedang jalan. Belum muncul di katalog dapat berasal dari akses atau data belum dimuat. Muat ulang tidak memulai discovery/profiling/sync.")
table(s, ["Pesan / status", "Arti praktis", "Tindakan"], [
    ["SOURCE_ACCESS_DENIED", "Backend tidak boleh membaca Sheet", "Cek akses Viewer service account"],
    ["SOURCE_READ_FAILED", "Google menolak tab / range / permintaan", "Cek sumber; Temukan tab dan profil ulang"],
    ["REVIEW_STALE", "Data sumber berubah sejak review", "Profiling dan review draft baru"],
    ["TYPE_OR_NULL_ERROR", "Tipe salah atau kolom wajib kosong", "Perbaiki nilai / parsing; dry-run ulang"],
    ["NEEDS_INPUT", "Batch menunggu verifikasi atau keputusan", "Buka Lihat detail; selesaikan temuan"],
    ["Belum muncul di katalog", "Periksa policy, role dan hasil pemuatan", "Uji akun pembaca; cek baris per tab"]
], [3.1, 4.1, 4.8], font=15)

# 24 — Catalog
s = base("Langkah 7 — Pastikan produk dapat dibaca", "Uji memakai akun pembaca sebenarnya; peran approver bukan akses otomatis ke semua data.", notes="Konfigurasi aktif dan policy disetujui merupakan gate berbeda. Pembaca harus memiliki assignment dan policy source/product yang tepat serta role yang diizinkan. Periksa sumber/tab dan waktu dimuat di katalog atau detail hasil ETL. Versi aktif lama dapat tetap terlihat selama revisi jika akses masih sah; baca freshness dan status terbaru.")
table(s, ["Checklist", "Bukti yang dilihat"], [
    ["Konfigurasi sudah ACTIVE", "Versi aktif pada produk / sumber"],
    ["Data sudah dimuat", "Baris dimuat dan waktu pemuatan per tab"],
    ["Akses sumber disetujui", "Metadata / policy dan scope organisasi sesuai"],
    ["Pembaca memiliki izin", "Role, assignment, policy produk, kolom dan baris"],
    ["Angka dapat ditelusuri", "Sumber / tab, versi, periode dan freshness"]
], [4.9, 7.1], font=19)
ribbon(s, "Muat ulang katalog menampilkan status terbaru; tidak menjalankan sync data.")

# 25 — Analysis question
s = base("Langkah 8 — Ajukan pertanyaan yang lengkap", "Sebut ukuran, kelompok, periode dan batas cakupan yang ingin dibandingkan.", section="BAGIAN 3 • ANALITIK, CHART DAN ACTION PLAN", notes="Gunakan pertanyaan contoh dan pastikan metriknya sudah tersedia. Query builder memberi kontrol eksplisit dimensi, metrik, filter, sort, serta join approved. AI dapat meminta klarifikasi, misalnya untuk periode ambigu. Query hasil data mengikuti izin pembaca. Jangan menuntut chart atas metrik yang belum didefinisikan.")
shape(s, 0.65, 2.33, 12, 1.26, LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.96, 2.59, 11.3, 0.88, "“Berapa jumlah keluhan per produk pada Januari 2026,\ndan berapa yang melewati SLA respons?”", 23, GREEN, True)
two(s, ("Pertanyaan AI", "Tulis pertanyaan di Dashboard.\nPilih ruang data bila diperlukan.\nJawab klarifikasi yang muncul."),
    ("Query builder", "Pilih metrik + dimensi.\nTetapkan filter periode / cakupan.\nPeriksa tabel hasil sebelum chart."), y=3.95, h=2.20)
ribbon(s, "Periksa definisi metrik, periode, jumlah record dan akses sebelum menyimpulkan hasil.")

# 26 — Native charts
s = base("Chart membuat pola lebih mudah dibaca", "Pilih bar untuk perbandingan dan line untuk tren pada periode yang sebanding.", source="USER_GUIDE_END_TO_END.md", notes="Dua chart ini ilustrasi editable PowerPoint, bukan screenshot atau data production. Januari 32 keluhan dan 12 pelanggaran, Februari 40 dan 8, Maret 36 dan 3. Rasio 37,5%, 20% dan 8,3%. Perubahan rasio adalah pola deskriptif; tidak membuktikan penyebab atau efek tindakan. Pilihan chart aplikasi mengikuti field/metrik hasil query.")
text(s, 0.78, 2.20, 5.6, 0.5, "Jumlah keluhan dan pelanggaran", 20, GREEN, True)
text(s, 6.88, 2.20, 5.5, 0.5, "Rasio respons melewati SLA", 20, GREEN, True)
data = CategoryChartData()
data.categories = ["Jan", "Feb", "Mar"]
data.add_series("Keluhan", [32, 40, 36])
data.add_series("Lewat SLA", [12, 8, 3])
chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(2.85), Inches(5.75), Inches(3.0), data).chart
data2 = CategoryChartData()
data2.categories = ["Jan", "Feb", "Mar"]
data2.add_series("Rasio (%)", [37.5, 20, 100*3/36])
chart2 = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(6.82), Inches(2.85), Inches(5.75), Inches(3.0), data2).chart
for c in (chart, chart2):
    c.has_title = False
    c.has_legend = c == chart
    if c.has_legend:
        c.legend.position = XL_LEGEND_POSITION.BOTTOM
        c.legend.font.name, c.legend.font.size = FONT, Pt(13)
        c.legend.include_in_layout = False
    c.chart_style = 10
    c.category_axis.tick_labels.font.name, c.category_axis.tick_labels.font.size = FONT, Pt(14)
    c.value_axis.tick_labels.font.name, c.value_axis.tick_labels.font.size = FONT, Pt(13)
    c.value_axis.minimum_scale = 0
    c.value_axis.maximum_scale = 50
    c.value_axis.major_unit = 10
    c.plots[0].has_data_labels = True
    c.plots[0].data_labels.font.name, c.plots[0].data_labels.font.size = FONT, Pt(13)
for i, series in enumerate(chart.series):
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb(GREEN if i == 0 else "D6A246")
    series.format.line.fill.background()
chart2.series[0].format.line.color.rgb = rgb(BLUE)
chart2.series[0].format.line.width = Pt(3)
chart2.plots[0].data_labels.position = XL_LABEL_POSITION.ABOVE
chart2.plots[0].data_labels.number_format = '0.0"%"'
chart2.value_axis.tick_labels.number_format = '0"%"'
ribbon(s, "SIMULASI • Periode dan metode harus sebanding; pola chart tidak membuktikan sebab-akibat.", GOLD)

# 27 — Five examples
s = base("Lima template untuk kebutuhan bisnis", "Setiap file memuat petunjuk, data contoh, master, formula dan aturan action plan.", source="CONTOH_DATA_ANALITIK.md", notes="Gunakan paket template yang tersedia di aplikasi. Seluruh data dan target fiktif. Marketing ATL tidak diberi angka atribusi palsu jika tidak terukur. CSAT bukan NPS. Competitiveness menormalkan ukuran kemasan dan membatasi keterbandingan. KPI memakai HIGHER/LOWER sesuai indikator. Master identik cukup dimuat sekali dan digunakan ulang.")
table(s, ["Template", "Pertanyaan analitik", "Usulan tindak lanjut contoh"], [
    ["Keluhan pelanggan", "Di mana keluhan dan pelanggaran SLA meningkat?", "Review antrean dan pembagian shift"],
    ["Marketing ATL / BTL", "Berapa biaya dan qualified leads teratribusi?", "Review campaign / biaya per lead"],
    ["Kepuasan per produk", "Berapa puas / respons valid pada metode setara?", "Verifikasi sampel dan masalah produk"],
    ["Competitiveness", "Bagaimana harga per unit dan skor layanan?", "Review layanan pada perbandingan setara"],
    ["Plan / goal / KPI", "Apa gap dan capaian menurut arah indikator?", "Review pipeline atau rencana pencapaian"]
], [3.0, 4.85, 4.15], font=16)
ribbon(s, "SIMULASI • Target / bobot contoh harus disesuaikan dan disetujui pemilik bisnis.", GOLD)

# 28 — Rules workflow
s = base("Langkah 9 — Aturan action plan dan formula", "Alur ini sudah tersedia di backend API, terpisah dari wizard ETL dan chat AI.", source="PANDUAN_ATURAN_ANALITIK.md", notes="Formula_Analitik memuat formula_code, expression, output_unit, description. Aturan_Action_Plan memuat rule_code, formula_code, operator, threshold, threshold_upper, action_description, owner_team_code, priority, due_days. Header harus sesuai template. Formula menggunakan kode metrik yang tersedia dan + - * / serta kurung, tanpa awalan =Excel. Aturan mengacu formula dan angka ambang. 50% pada RATIO = 0.50. Impor dari sumber yang sama dengan produk, preview, reviewer sumber berbeda approves hash, lalu evaluate. Pengeditan tab/semantic memerlukan draft dan review baru; REVOKE menghentikan aturan approved.")
two(s, ("Formula_Analitik", "formula_code: sla_breach_ratio\nRumus: total pelanggaran / total keluhan\noutput_unit: RATIO\n\nGunakan kode metrik produk aktif.\nTab kontrol: Dikecualikan dari ETL."),
    ("Aturan_Action_Plan", "formula_code: sla_breach_ratio\noperator: gt • threshold: 0.10\nAction: review antrean dan shift\nOwner: TEAM-CS • tenggat: 7 hari\n\n50% ditulis 0.50, bukan 50."))
ribbon(s, "API: impor draft → preview → approval reviewer berbeda → evaluate; REVOKE untuk menghentikan.", GOLD)

# 29 — Action output
s = base("Contoh output: usulan dengan bukti angka", "Kondisi cocok menghasilkan PROPOSED; penyebab tetap hipotesis sampai diverifikasi.", source="PANDUAN_ATURAN_ANALITIK.md", notes="Ilustrasi Januari memiliki 32 keluhan dan 12 pelanggaran. Formula 12/32 = 0.375, lebih besar dari threshold 0.10. Evaluasi memberi usulan review antrean/shift, tim TEAM-CS, tenggat tujuh hari dari tanggal evaluasi UTC dan cause_status UNVERIFIED. Backend tidak menjalankan tindakan, menulis Sheet, mengirim notifikasi atau memberi hak akses berdasarkan owner_team_code. NULL atau denominator nol memberi blocker dan tidak memicu aturan terkait.")
for i, (title, body) in enumerate([
    ("BUKTI", "Periode: Januari 2026\n32 keluhan • 12 lewat SLA\n12 / 32 = 37,5%\n\nAturan: lebih besar dari 10%"),
    ("USULAN", "Verifikasi antrean dan shift.\nUsulkan perbaikan SLA setelah\nbukti operasional diperiksa.\n\nStatus: PROPOSED"),
    ("TANGGUNG JAWAB", "Owner: TEAM-CS\nTenggat: 7 hari sejak evaluasi\nPenyebab: UNVERIFIED\n\nRencana ditinjau pemilik bisnis")]):
    card(s, 0.65 + i * 4.08, 2.38, 3.85, 3.68, title, body, size=18)
ribbon(s, "SIMULASI • NULL / penyebut nol memberi blocker; tidak memicu tindakan terkait.", GOLD)

# 30 — Follow-up
s = base("Langkah 10 — Laksanakan dan ukur ulang", "Bedakan usulan, keputusan bisnis, pelaksanaan dan bukti hasil.", source="PANDUAN_FORMAT_GOOGLE_SHEET.md", notes="Pemilik bisnis menentukan apakah usulan masuk akal dan menambah rencana pelaksanaan yang terukur. Catat baseline, target, owner, due_date, completed_at, periode follow-up, followup_value dan evidence_ref. Bandingkan metode serta populasi yang sebanding dan catat intervensi lain. Backend aturan hanya menghitung usulan; pemantauan pelaksanaan saat ini dicatat manusia, misalnya melalui tab Action_Plan. Kosong bukan hasil nol.")
table(s, ["Tahap", "Yang harus dicatat", "Contoh hasil ideal"], [
    ["Verifikasi masalah", "Bukti, periode, baseline, metode dan sampel", "Masalah dan batas interpretasi jelas"],
    ["Putuskan tindakan", "Target, pemilik, tenggat dan cara ukur", "Rencana spesifik yang disetujui bisnis"],
    ["Laksanakan", "Status, waktu selesai dan bukti aktivitas", "Pelaksanaan dapat ditelusuri"],
    ["Ukur follow-up", "Periode baru, hasil dan konteks pembanding", "Angka dihitung dengan metode sebanding"],
    ["Evaluasi", "Verifikasi hasil, faktor lain dan langkah berikutnya", "Keputusan lanjut / revisi berbasis bukti"]
], [2.45, 5.15, 4.4], font=16)
ribbon(s, "Belum dievaluasi → follow-up kosong / NOT_EVALUATED; tidak dianggap berhasil atau bernilai 0.")

# 31 — Ideal acceptance
s = base("Hasil ideal yang ingin dicapai", "Tetapkan bukti keberhasilan bersama pemilik bisnis sebelum memperluas penggunaan.", source="PANDUAN_FORMAT_GOOGLE_SHEET.md", notes="Jadikan checklist ini acceptance criteria untuk pilot. Hindari janji angka peningkatan yang belum diukur. Target manfaat ditetapkan bisnis dan dievaluasi pada periode follow-up. Kepercayaan analitik tidak cukup dengan job sukses; kualitas, akses, definisi dan cakupan juga harus benar.")
table(s, ["Sasaran", "Bukti keberhasilan"], [
    ["Data konsisten dan berkualitas", "Key / tipe / referensi valid; temuan ditangani"],
    ["Waktu dan konteks jelas", "Periode, metode, satuan, sampel dan freshness tersedia"],
    ["Analitik dapat diulang", "Definisi metrik dan filter menghasilkan angka yang dapat diperiksa"],
    ["Chart membantu keputusan", "Kelompok / periode setara dan visual tidak menyesatkan"],
    ["Usulan dapat ditindaklanjuti", "Aturan, bukti, pemilik, target dan tenggat jelas"],
    ["Manfaat terverifikasi", "Baseline dan follow-up diperiksa, bukan klaim otomatis"]
], [4.5, 7.5], font=17)

# 32 — Training close / sources
s = base("Mulai dengan satu dataset dan satu pertanyaan", "Latihan: keluhan pelanggan → SLA per produk → usulan perbaikan → evaluasi.", source="CONTOH_DATA_ANALITIK.md", notes="Demo menggunakan file Keluhan Pelanggan dalam paket. Urutan demo: cek BACA_DULU/Kamus_Data/Rencana_ETL, konversi ke Sheets, hubungkan dan profiling, klasifikasikan tab, isi tujuh bagian konfigurasi, review/deploy, sync dan validasi per tab, buka katalog dan query/chart. Demonstrasikan aturan action plan lewat API terpisah bila peserta teknis tersedia. Jangan mengubah sumber production aktif dengan data contoh. Tautan dalam slide dapat diklik. Catatan pembicara setiap slide memuat referensi dan penjelasan tambahan.")
card(s, 0.65, 2.29, 7.32, 3.78, "Checklist demo", "1. Baca format dan arti kolom pada template.\n2. Hubungkan, klasifikasikan dan setup config.\n3. Review → deploy → sync → cek katalog.\n4. Query metrik dan tampilkan chart.\n5. Review aturan API dan periksa usulan.\n6. Sepakati baseline, target dan ukur ulang.", size=19)
card(s, 8.23, 2.29, 4.37, 3.78, "Bahan pendukung", "Panduan aplikasi, lima template\ndan panduan aturan tersedia\nmelalui tautan di bawah.\n\nSemua contoh memakai data fiktif.", size=19)
link_button(s, 0.65, 6.22, 3.86, "Buka panduan aplikasi", APP + "/guide")
link_button(s, 4.70, 6.22, 3.86, "Unduh 5 template Excel", TEMPLATES)
link_button(s, 8.75, 6.22, 3.86, "Panduan aturan / formula", RULES_GUIDE)

prs.save(FILE)
manifest[0]["title"] = "Dari Google Sheets ke Analitik dan Action Plan"
# PowerPoint applies its hyperlink theme even when a run has an explicit color.
# Links here sit on dark green buttons, so use white for both hyperlink states.
import io
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

with ZipFile(FILE) as archive:
    entries = {name: archive.read(name) for name in archive.namelist()}
ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
theme = ET.fromstring(entries["ppt/theme/theme1.xml"])
for kind in ("hlink", "folHlink"):
    scheme = theme.find(f"a:themeElements/a:clrScheme/a:{kind}", ns)
    for child in list(scheme):
        scheme.remove(child)
    ET.SubElement(scheme, "{" + ns["a"] + "}srgbClr", {"val": WHITE})
entries["ppt/theme/theme1.xml"] = ET.tostring(theme, encoding="utf-8", xml_declaration=True)
stream = io.BytesIO()
with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
    for name, content in entries.items():
        archive.writestr(name, content)
FILE.write_bytes(stream.getvalue())
import json

(OUT / "slide_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"Created {len(prs.slides)} editable slides: {FILE}")
